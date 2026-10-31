from __future__ import annotations

import base64
import copy
import hashlib
import io
import json
import re
import uuid
import zipfile
from collections.abc import Awaitable, Callable, Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from typing import Literal, cast
from uuid import UUID

from pptx import Presentation
from sqlalchemy import delete, or_, select, text
from sqlalchemy.engine import URL, make_url
from sqlalchemy.exc import ArgumentError
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from revenueos.business_case_contracts import (
    BusinessCaseApprovalRequest,
    BusinessCaseCalculateRequest,
    BusinessCaseCreateRequest,
    BusinessCaseInputValue,
    BusinessCaseScenarioRequest,
    ScenarioCalculationResponse,
    ScenarioOverride,
    SensitivityRequest,
    ValueModelApprovalRequest,
    ValueModelCreateRequest,
    ValueModelDefinition,
)
from revenueos.business_case_services import BusinessCaseService
from revenueos.config import Settings, get_settings
from revenueos.create_contracts import (
    PresentationApprovalRequest,
    PresentationAudienceInput,
    PresentationBriefRequest,
    PresentationClaimDecision,
    PresentationGenerateRequest,
    PresentationReviewRequest,
    TemplateApprovalRequest,
    TemplateSlideUpdate,
    TemplateSummaryResponse,
    TemplateUploadRequest,
)
from revenueos.create_pptx import PPTX_MIME_TYPE, PptxProcessingError
from revenueos.create_services import CreateService
from revenueos.create_worker import (
    ClaimedCreateWork,
    CreateWorkerService,
    _compose,
    _output_expectations,
    _render_manifest,
    _validate_claim_manifest,
    create_processor,
)
from revenueos.database import set_tenant_database_context
from revenueos.deal_room_contracts import (
    DealRoomDraftContent,
    DealRoomDraftUpdate,
    DealRoomMilestoneDraft,
    DealRoomPublishRequest,
    DealRoomResourceDraft,
    DealRoomStakeholderDraft,
)
from revenueos.deal_room_services import DealRoomService
from revenueos.demo_data import (
    DEMO_DOCUMENTS,
    DEMO_NAMESPACE,
    DEMO_TARGET_MARKET_NAME,
    demo_campaign_ids,
    demo_companion_ids,
    demo_debrief_ids,
    demo_ids,
    demo_interaction_ids,
    demo_live_ids,
    demo_marker_ids,
    demo_phone_contact_id,
    demo_sales_analytics_ids,
    demo_sales_forecast_ids,
    demo_sales_target_ids,
    demo_source_evidence_ids,
    demo_visual_ids,
    seed_demo_data,
)
from revenueos.domain import (
    DealRoomMilestoneStatus,
    DealRoomOwnerParty,
    DealRoomResourceKind,
    HandoverAuthorityType,
    HandoverSourceType,
)
from revenueos.handover_contracts import SECTION_KEYS, HandoverContent
from revenueos.handover_services import HandoverService, SourceDefinition
from revenueos.models import (
    Base,
    BetaSystemEvent,
    ClosedWonHandover,
    ClosedWonHandoverAuditEvent,
    ClosedWonHandoverRevision,
    ClosedWonHandoverSource,
    Company,
    CompanyKnowledgeSource,
    Contact,
    CreateApprovedContentItem,
    CreateBusinessCase,
    CreateBusinessCaseVersion,
    CreatePresentation,
    CreatePresentationVersion,
    CreateTemplate,
    CreateTemplateSlide,
    CreateTemplateVersion,
    CreateValueModel,
    CreateValueModelVersion,
    DealRoom,
    DealRoomAccessLink,
    DealRoomAuditEvent,
    DealRoomRevision,
    EngageCampaign,
    EngageCampaignAudience,
    Opportunity,
    Organisation,
    OrganisationMembership,
    OrganisationMethodologySetting,
    ProspectTargetMarket,
    ProspectTargetMarketVersion,
    SalesForecastJudgment,
    SalesForecastJudgmentRevision,
    SalesForecastPeriod,
    SalesForecastReviewerJudgment,
    SalesForecastReviewerRevision,
    SalesPipeline,
    SalesPipelineStage,
    SalesTarget,
    SalesTargetRevision,
    SellingProfile,
    SellingProfileRevision,
    Task,
)
from revenueos.selling_profile_contracts import (
    CompanyKnowledgeSourceCreate,
    SellingProfileApproveRequest,
    SellingProfileContent,
    SellingProfileDraftCreate,
)
from revenueos.selling_profile_services import SellingProfileService
from revenueos.tenant import Role, TenantContext
from revenueos.visual_storage import VisualStorage, VisualStorageError, create_visual_storage

OWNER_DEMO_DATASET_VERSION = 22
OWNER_PROFILE_KEY = "owner-demo-selling-profile-v1"
OWNER_KNOWLEDGE_WEBSITE_KEY = "owner-demo-knowledge-website-v1"
OWNER_KNOWLEDGE_TEXT_KEY = "owner-demo-knowledge-text-v1"
OWNER_VALUE_MODEL_KEY = "owner-demo-value-model-v1"
OWNER_BUSINESS_CASE_KEY = "owner-demo-business-case-v1"
OWNER_BUSINESS_CASE_CALCULATION_KEY = "owner-demo-business-case-calculation-v1"
OWNER_VALUE_MODEL_NAME = "[DEMO] Relationship workflow value model"
OWNER_VALUE_MODEL_DESCRIPTION = "A bounded synthetic value model for the private owner demonstration."
OWNER_BUSINESS_CASE_TITLE = "[DEMO] Northstar workflow Business Case"
OWNER_TEMPLATE_NAME = "[DEMO] Oryntela customer story"
OWNER_TEMPLATE_FILE_NAME = "oryntela-owner-demo.pptx"
OWNER_PRESENTATION_KEY = "owner-demo-presentation-v1"
OWNER_PRESENTATION_GENERATION_KEY = "owner-demo-presentation-generation-v1"
OWNER_PRESENTATION_TITLE = "[DEMO] Northstar workflow pilot presentation"
OWNER_PRESENTATION_FOCUS = "Keep the synthetic security, implementation and customer-owned next steps concise."
OWNER_RESET_CLEANUP_EVENT = "owner_demo_reset_cleanup_pending"
OWNER_RESET_CLEANUP_SCHEMA_VERSION = 1
OWNER_RESET_CLEANUP_MAX_KEYS = 16
OWNER_BASE_RESERVATION_EVENT = "owner_demo_base_reservation"
OWNER_BASE_MANIFEST_EVENT = "owner_demo_base_manifest"
OWNER_BASE_MANIFEST_SCHEMA_VERSION = 1
BASE_DEMO_DATASET_VERSION = 18
OWNER_EXTENSION_MANIFEST_KEYS = frozenset(
    {
        "selling_profile_revision_id",
        "knowledge_website_source_id",
        "knowledge_text_source_id",
        "value_model_id",
        "business_case_id",
        "business_case_version_id",
        "template_version_id",
        "presentation_id",
        "presentation_version_id",
        "deal_room_id",
        "handover_id",
        "handover_revision_id",
        "task_id",
    }
)


@dataclass(frozen=True)
class _MethodologySettingSnapshot:
    selection: str
    custom_definition_id: UUID | None
    updated_by_user_id: UUID
    created_at: datetime
    updated_at: datetime


@dataclass(frozen=True)
class _MethodologyReservation:
    preexisting: bool
    original: _MethodologySettingSnapshot | None


def owner_demo_template_id(organisation_id: UUID) -> UUID:
    """Return the tenant-scoped identity marker for the owner-demo template."""

    return uuid.uuid5(DEMO_NAMESPACE, f"{organisation_id}:owner-demo-template-v1")


def owner_demo_cleanup_event_id(organisation_id: UUID) -> UUID:
    """Return the durable exact-key cleanup marker for an owner-demo reset."""

    return uuid.uuid5(DEMO_NAMESPACE, f"{organisation_id}:owner-demo-reset-cleanup-v1")


def owner_demo_base_manifest_id(organisation_id: UUID) -> UUID:
    """Return the durable marker proving the owner command created the base fixture."""

    return uuid.uuid5(DEMO_NAMESPACE, f"{organisation_id}:owner-demo-base-manifest-v1")


def owner_demo_task_id(organisation_id: UUID) -> UUID:
    """Return the tenant-scoped identity for the canonical owner-demo Task."""

    return uuid.uuid5(DEMO_NAMESPACE, f"{organisation_id}:owner-demo-task-v1")


def _validated_local_database_url(raw_url: str | URL | None) -> URL:
    try:
        if raw_url is None:
            raise ValueError
        database_url = make_url(raw_url)
    except (ArgumentError, TypeError, ValueError):
        raise ValueError("database must be local SQLite or loopback PostgreSQL") from None
    driver = database_url.drivername.split("+", 1)[0]
    query_keys = {str(key).casefold() for key in database_url.query}
    if driver == "sqlite":
        database = database_url.database
        if (
            database_url.host is not None
            or database is None
            or not database
            or database.startswith("file:")
            or "://" in database
            or database_url.query
        ):
            raise ValueError("database must be local SQLite or loopback PostgreSQL")
    elif driver == "postgresql":
        if database_url.host not in {"localhost", "127.0.0.1", "::1"} or query_keys.intersection(
            {"host", "hostaddr", "unix_sock", "unix_socket"}
        ):
            raise ValueError("database must be local SQLite or loopback PostgreSQL")
    else:
        raise ValueError("database must be local SQLite or loopback PostgreSQL")
    return database_url


def _assert_private_demo_settings(settings: Settings) -> None:
    unsafe_reasons: list[str] = []
    if settings.environment not in {"development", "test"}:
        unsafe_reasons.append("environment must be development or test")
    if settings.auth_mode != "mock" or not settings.mock_auth_enabled:
        unsafe_reasons.append("mock authentication must be enabled")
    if settings.visual_storage_backend != "local":
        unsafe_reasons.append("visual storage must be local")
    if settings.ai_provider_name != "mock" or settings.feature_openai_provider_enabled:
        unsafe_reasons.append("AI must use the deterministic mock provider")
    if settings.prospect_research_provider_name != "mock" or settings.feature_prospect_external_provider_enabled:
        unsafe_reasons.append("Prospect must use the deterministic mock provider")
    if any(
        (
            settings.feature_microsoft_365_enabled,
            settings.feature_google_workspace_enabled,
            settings.feature_hubspot_crm_enabled,
            settings.feature_salesforce_crm_enabled,
        )
    ):
        unsafe_reasons.append("Microsoft 365, Google Workspace, HubSpot and Salesforce must remain disabled")
    if settings.private_beta_real_data_enabled:
        unsafe_reasons.append("real-customer data release must remain disabled")
    if settings.billing_mode != "test" or settings.billing_provider_name != "deterministic":
        unsafe_reasons.append("billing must remain on the deterministic test adapter")
    if settings.feature_billing_enabled or settings.feature_credits_enabled:
        unsafe_reasons.append("billing and Credits must remain disabled")
    try:
        _validated_local_database_url(settings.database_url)
    except ValueError as exc:
        unsafe_reasons.append(str(exc))
    if unsafe_reasons:
        raise ValueError("Owner demo refused: " + "; ".join(unsafe_reasons) + ".")


def _assert_private_demo_database_binding(
    settings: Settings,
    session_factory: async_sessionmaker[AsyncSession],
) -> None:
    expected = _validated_local_database_url(settings.database_url)
    bind = session_factory.kw.get("bind")
    bound_url = getattr(bind, "url", None)
    try:
        actual = _validated_local_database_url(bound_url)
    except ValueError as exc:
        raise ValueError(f"Owner demo refused: session factory {exc}.") from None
    if actual != expected:
        raise ValueError("Owner demo refused: session factory database does not match the validated local database.")


def _base_fixture_collision(resource: str) -> ValueError:
    return ValueError(f"Owner demo refused: the deterministic base {resource} is not the exact synthetic fixture.")


def _reserved_base_demo_ids(organisation_id: UUID) -> frozenset[UUID]:
    """Return every fixed identity the base seed may adopt instead of creating."""

    company_id, opportunity_id, meeting_ids, transcript_ids = demo_ids(organisation_id)
    analytics_opportunity_ids, analytics_interaction_ids = demo_sales_analytics_ids(organisation_id)
    companion_groups = demo_companion_ids(organisation_id)
    reserved = {
        company_id,
        opportunity_id,
        *meeting_ids,
        *transcript_ids,
        demo_phone_contact_id(organisation_id),
        *demo_campaign_ids(organisation_id).values(),
        *analytics_opportunity_ids,
        *analytics_interaction_ids,
        *demo_sales_target_ids(organisation_id).values(),
        *demo_sales_forecast_ids(organisation_id).values(),
        *demo_source_evidence_ids(organisation_id).values(),
        *demo_interaction_ids(organisation_id),
        *demo_debrief_ids(organisation_id),
        *demo_visual_ids(organisation_id),
        *demo_marker_ids(organisation_id),
        *demo_live_ids(organisation_id).values(),
    }
    for group in companion_groups:
        reserved.update(group)
    labels = (
        "transcript-version-1",
        "transcript-version-2",
        "online-metadata-teams",
        "online-metadata-zoom",
        "online-transcript-import",
        "online-transcript-capture",
        "online-transcript-evidence",
        "online-recording-import",
        "online-recording-capture",
        "online-recording-evidence",
        "online-recording-consent",
        "online-google-debrief",
        "online-google-meeting",
        "online-metadata-google",
        "online-google-debrief-session",
        "daily-action:send-security-summary",
        "daily-action:confirm-pilot-scope",
        "methodology-final-intelligence",
        "methodology-brain-snapshot",
        "methodology-final-evidence",
    )
    reserved.update(uuid.uuid5(DEMO_NAMESPACE, f"{organisation_id}:{label}") for label in labels)
    reserved.update(
        uuid.uuid5(DEMO_NAMESPACE, f"{organisation_id}:sales-analytics-stage-{index:02d}-{step:02d}")
        for index in range(20)
        for step in range(6)
    )
    return frozenset(reserved)


async def _assert_no_reserved_base_identity_collisions(
    session: AsyncSession,
    organisation_id: UUID,
) -> None:
    """Reject any pre-existing base identity before the ownership manifest exists."""

    reserved_ids = tuple(_reserved_base_demo_ids(organisation_id))
    for table in Base.metadata.tables.values():
        if "id" not in table.c or "organisation_id" not in table.c:
            continue
        collision = await session.scalar(
            select(table.c.id)
            .where(
                table.c.organisation_id == organisation_id,
                table.c.id.in_(reserved_ids),
            )
            .limit(1)
        )
        if collision is not None:
            raise _base_fixture_collision(f"reserved {table.name} identity")


async def _assert_base_demo_fixture_safe(
    session_factory: async_sessionmaker[AsyncSession],
    organisation_id: UUID,
    expected_user_id: UUID | None = None,
    *,
    allow_owner_lifecycle: bool = False,
) -> None:
    """Prove ownership of base fixture roots before seed or reset can mutate them.

    The base seed predates the stricter owner-demo marker.  Its durable
    ``demo_data_seeded`` event is therefore the authority marker, while exact
    root checks prevent a later row at a reserved UUID from being adopted or
    deleted as demo data.
    """

    company_id, opportunity_id, _, _ = demo_ids(organisation_id)
    phone_contact_id = demo_phone_contact_id(organisation_id)
    campaign_ids = demo_campaign_ids(organisation_id)
    analytics_opportunity_ids, _ = demo_sales_analytics_ids(organisation_id)
    target_ids = demo_sales_target_ids(organisation_id)
    forecast_ids = demo_sales_forecast_ids(organisation_id)
    async with session_factory() as session:
        await set_tenant_database_context(session, organisation_id)
        organisation = await session.get(Organisation, organisation_id)
        company = await session.scalar(
            select(Company).where(Company.organisation_id == organisation_id, Company.id == company_id)
        )
        opportunity = await session.scalar(
            select(Opportunity).where(
                Opportunity.organisation_id == organisation_id,
                Opportunity.id == opportunity_id,
            )
        )
        phone_contact = await session.scalar(
            select(Contact).where(Contact.organisation_id == organisation_id, Contact.id == phone_contact_id)
        )
        owner_task = await session.scalar(
            select(Task).where(
                Task.organisation_id == organisation_id,
                Task.id == owner_demo_task_id(organisation_id),
            )
        )
        legacy_company = await session.scalar(
            select(Company).where(
                Company.organisation_id == organisation_id,
                Company.id == campaign_ids["company"],
            )
        )
        campaign_contacts = list(
            (
                await session.scalars(
                    select(Contact).where(
                        Contact.organisation_id == organisation_id,
                        Contact.id.in_((campaign_ids["contact-jane"], campaign_ids["contact-sam"])),
                    )
                )
            ).all()
        )
        campaign = await session.scalar(
            select(EngageCampaign).where(
                EngageCampaign.organisation_id == organisation_id,
                EngageCampaign.id == campaign_ids["campaign"],
            )
        )
        campaign_audience_company_ids = set(
            (
                await session.scalars(
                    select(EngageCampaignAudience.company_id).where(
                        EngageCampaignAudience.organisation_id == organisation_id,
                    )
                )
            ).all()
        )
        analytics_opportunities = list(
            (
                await session.scalars(
                    select(Opportunity).where(
                        Opportunity.organisation_id == organisation_id,
                        Opportunity.id.in_(analytics_opportunity_ids),
                    )
                )
            ).all()
        )
        target_markets = list(
            (
                await session.scalars(
                    select(ProspectTargetMarket).where(
                        ProspectTargetMarket.organisation_id == organisation_id,
                        ProspectTargetMarket.name == DEMO_TARGET_MARKET_NAME,
                    )
                )
            ).all()
        )
        sales_targets = list(
            (
                await session.scalars(
                    select(SalesTarget).where(
                        SalesTarget.organisation_id == organisation_id,
                        SalesTarget.id.in_(
                            (target_ids["won-value"], target_ids["meetings"], target_ids["quarter-won-value"])
                        ),
                    )
                )
            ).all()
        )
        sales_target_revisions = list(
            (
                await session.scalars(
                    select(SalesTargetRevision).where(
                        SalesTargetRevision.organisation_id == organisation_id,
                        SalesTargetRevision.id.in_(
                            (
                                target_ids["won-value-revision"],
                                target_ids["meetings-revision"],
                                target_ids["quarter-won-value-revision"],
                            )
                        ),
                    )
                )
            ).all()
        )
        forecast_period = await session.scalar(
            select(SalesForecastPeriod).where(
                SalesForecastPeriod.organisation_id == organisation_id,
                SalesForecastPeriod.id == forecast_ids["period"],
            )
        )
        forecast_judgments = list(
            (
                await session.scalars(
                    select(SalesForecastJudgment).where(
                        SalesForecastJudgment.organisation_id == organisation_id,
                        SalesForecastJudgment.id.in_(
                            (
                                forecast_ids["main-judgment"],
                                forecast_ids["likely-judgment"],
                                forecast_ids["possible-judgment"],
                                forecast_ids["excluded-judgment"],
                            )
                        ),
                    )
                )
            ).all()
        )
        forecast_revisions = list(
            (
                await session.scalars(
                    select(SalesForecastJudgmentRevision).where(
                        SalesForecastJudgmentRevision.organisation_id == organisation_id,
                        SalesForecastJudgmentRevision.id.in_(
                            (
                                forecast_ids["main-revision"],
                                forecast_ids["likely-revision"],
                                forecast_ids["possible-revision"],
                                forecast_ids["excluded-revision"],
                            )
                        ),
                    )
                )
            ).all()
        )
        reviewer_judgments = list(
            (
                await session.scalars(
                    select(SalesForecastReviewerJudgment).where(
                        SalesForecastReviewerJudgment.organisation_id == organisation_id,
                        SalesForecastReviewerJudgment.id.in_(
                            (
                                forecast_ids["main-reviewer-judgment"],
                                forecast_ids["likely-reviewer-judgment"],
                                forecast_ids["possible-reviewer-judgment"],
                            )
                        ),
                    )
                )
            ).all()
        )
        reviewer_revisions = list(
            (
                await session.scalars(
                    select(SalesForecastReviewerRevision).where(
                        SalesForecastReviewerRevision.organisation_id == organisation_id,
                        SalesForecastReviewerRevision.id.in_(
                            (
                                forecast_ids["main-reviewer-revision"],
                                forecast_ids["likely-reviewer-revision"],
                                forecast_ids["possible-reviewer-revision"],
                            )
                        ),
                    )
                )
            ).all()
        )
        markers = list(
            (
                await session.scalars(
                    select(BetaSystemEvent).where(
                        BetaSystemEvent.organisation_id == organisation_id,
                        BetaSystemEvent.event_type == "demo_data_seeded",
                        BetaSystemEvent.subject_id == opportunity_id,
                    )
                )
            ).all()
        )
        owner_manifest_event = await session.scalar(
            select(BetaSystemEvent).where(
                BetaSystemEvent.organisation_id == organisation_id,
                BetaSystemEvent.id == owner_demo_base_manifest_id(organisation_id),
            )
        )
        methodology_setting = await session.get(
            OrganisationMethodologySetting,
            organisation_id,
        )

        if owner_manifest_event is None and not markers:
            await _assert_no_reserved_base_identity_collisions(session, organisation_id)

        fixture_rows_exist = any(
            (
                company is not None,
                opportunity is not None,
                phone_contact is not None,
                owner_task is not None,
                legacy_company is not None,
                bool(campaign_contacts),
                campaign is not None,
                bool(analytics_opportunities),
                bool(target_markets),
                bool(sales_targets),
                bool(sales_target_revisions),
                forecast_period is not None,
                bool(forecast_judgments),
                bool(forecast_revisions),
                bool(reviewer_judgments),
                bool(reviewer_revisions),
            )
        )
        owner_manifest_base: dict[str, object] | None = None
        owner_extensions: dict[str, UUID] | None = None
        owner_has_reservation = False
        if owner_manifest_event is not None:
            manifest_user_id = expected_user_id or owner_manifest_event.actor_user_id
            if manifest_user_id is None:
                raise _base_fixture_collision("owner manifest")
            if owner_manifest_event.event_type == OWNER_BASE_RESERVATION_EVENT:
                methodology_reservation = _validated_base_reservation(
                    owner_manifest_event,
                    organisation_id,
                    manifest_user_id,
                )
                _assert_recoverable_methodology_reservation_state(
                    methodology_setting,
                    manifest_user_id,
                    methodology_reservation,
                )
                owner_has_reservation = True
            else:
                owner_manifest_base = _validated_base_manifest(
                    owner_manifest_event,
                    organisation_id,
                    manifest_user_id,
                )
                _validate_completed_owner_methodology_setting(
                    methodology_setting,
                    manifest_user_id,
                    owner_manifest_base,
                    allow_missing=expected_user_id is None,
                )
                owner_extensions = _validated_extension_manifest(
                    owner_manifest_event,
                    organisation_id,
                    manifest_user_id,
                )
        elif expected_user_id is not None and (markers or fixture_rows_exist):
            # The owner command never adopts a legacy base seed.  Only a clean
            # reservation or its own durable manifest can authorise mutation.
            raise _base_fixture_collision("owner manifest")

        if not markers:
            if fixture_rows_exist:
                raise _base_fixture_collision("ownership marker")
            if owner_manifest_event is not None and not owner_has_reservation:
                raise _base_fixture_collision("owner manifest")
            return
        if len(markers) != 1:
            raise _base_fixture_collision("ownership marker")
        marker = markers[0]
        owner_user_id = expected_user_id or marker.actor_user_id
        if (
            owner_user_id is None
            or marker.actor_user_id != owner_user_id
            or marker.metadata_json != {"dataset_version": BASE_DEMO_DATASET_VERSION}
        ):
            raise _base_fixture_collision("ownership marker")
        if company is None or opportunity is None or phone_contact is None:
            raise _base_fixture_collision("root set")
        if (
            company.name != "[DEMO] Northstar Facilities Group"
            or company.website != "https://northstar-facilities.example/"
            or company.normalized_domain != "northstar-facilities.example"
            or company.industry != "Synthetic multi-site facilities"
            or company.employee_count != 850
            or company.status != "prospect"
            or company.owner_user_id != owner_user_id
            or company.archived_at is not None
        ):
            raise _base_fixture_collision("Company")
        common_opportunity_is_exact = (
            opportunity.company_id != company_id
            or opportunity.name != "[DEMO] Northstar multi-site workflow pilot"
            or opportunity.currency != "AUD"
            or opportunity.owner_user_id != owner_user_id
            or opportunity.description != "Synthetic private-beta opportunity. No real customer or personal data."
            or opportunity.pipeline_id is None
            or opportunity.pipeline_stage_id is None
            or opportunity.archived_at is not None
        )
        pristine_opportunity = (
            opportunity.stage == "evaluation"
            and opportunity.status == "open"
            and opportunity.estimated_value == Decimal("420000.00")
            and opportunity.actual_close_date is None
            and opportunity.outcome_reason is None
            and opportunity.outcome_note is None
            and opportunity.outcome_provenance is None
        )
        owner_lifecycle_opportunity = (
            allow_owner_lifecycle
            and owner_extensions is not None
            and opportunity.stage == "closed_won"
            and opportunity.status == "won"
            and opportunity.estimated_value is not None
            and opportunity.estimated_value >= 0
            and opportunity.actual_close_date is not None
            and opportunity.outcome_reason is not None
            and opportunity.outcome_provenance == "seller_reported"
        )
        if common_opportunity_is_exact or not (pristine_opportunity or owner_lifecycle_opportunity):
            raise _base_fixture_collision("Opportunity")
        if (
            phone_contact.company_id != company_id
            or phone_contact.first_name != "[DEMO] Jordan"
            or phone_contact.last_name != "Lee"
            or phone_contact.email != "jordan.lee@example.test"
            or phone_contact.phone != "+61 400 000 017"
            or phone_contact.job_title != "Finance approver"
            or phone_contact.status != "active"
            or phone_contact.owner_user_id != owner_user_id
            or phone_contact.archived_at is not None
        ):
            raise _base_fixture_collision("Contact")
        if owner_task is not None:
            owner_task_due = owner_task.due_at
            if owner_task_due is not None and owner_task_due.utcoffset() is None:
                owner_task_due = owner_task_due.replace(tzinfo=UTC)
            if (
                owner_manifest_event is None
                or (owner_extensions is not None and owner_extensions["task_id"] != owner_task.id)
                or owner_task.company_id != company_id
                or owner_task.contact_id != phone_contact_id
                or owner_task.opportunity_id != opportunity_id
                or owner_task.title != "[DEMO] Prepare the Northstar security workshop"
                or owner_task.description
                != "Synthetic owner-demo follow-up. Confirm the reviewed security summary and workshop agenda."
                or owner_task.status not in {"open", "in_progress", "completed", "cancelled"}
                or owner_task.priority != "high"
                or owner_task_due != datetime(2026, 11, 4, 2, 0, tzinfo=UTC)
                or owner_task.assigned_user_id != owner_user_id
                or owner_task.created_by_user_id != owner_user_id
            ):
                raise _base_fixture_collision("Task")
        if legacy_company is not None and (
            legacy_company.name != "[DEMO] Northstar Facilities Group"
            or legacy_company.website != "https://northstar.example.test"
            or legacy_company.normalized_domain is not None
            or legacy_company.industry != "Synthetic multi-site facilities"
            or legacy_company.employee_count != 850
            or legacy_company.status != "prospect"
            or legacy_company.owner_user_id != owner_user_id
            or legacy_company.archived_at is not None
        ):
            raise _base_fixture_collision("legacy campaign Company")
        expected_campaign_contact_company = legacy_company.id if legacy_company is not None else company_id
        expected_contacts = {
            campaign_ids["contact-jane"]: (
                "[DEMO] Jane",
                "Smith",
                "jane.smith@northstar.example.test",
                "Chief Information Officer",
            ),
            campaign_ids["contact-sam"]: (
                "[DEMO] Sam",
                "Rivera",
                None,
                "Technology Director",
            ),
        }
        for contact in campaign_contacts:
            expected = expected_contacts.get(contact.id)
            if expected is None or (
                contact.company_id != expected_campaign_contact_company
                or (contact.first_name, contact.last_name, contact.email, contact.job_title) != expected
                or contact.phone is not None
                or contact.status != "active"
                or contact.owner_user_id != owner_user_id
                or contact.archived_at is not None
            ):
                raise _base_fixture_collision("campaign Contact")
        if campaign is not None and (
            campaign.owner_user_id != owner_user_id or campaign.state != "paused" or campaign.current_version != 1
        ):
            raise _base_fixture_collision("Campaign")
        if (
            expected_user_id is not None
            and legacy_company is not None
            and legacy_company.id in campaign_audience_company_ids
        ):
            raise ValueError(
                "Owner demo dataset upgrade required: the published campaign audience belongs to a legacy "
                "Account. Run pnpm demo:reset, then pnpm demo:owner."
            )
        analytics_by_id = {item.id: item for item in analytics_opportunities}
        for index, analytics_id in enumerate(analytics_opportunity_ids):
            item = analytics_by_id.get(analytics_id)
            if item is not None and (
                item.company_id != company_id
                or item.name != f"[DEMO] Analytics opportunity {index + 1:02d}"
                or item.owner_user_id != owner_user_id
                or item.description != "Synthetic analytics history. No real customer or personal data."
                or item.archived_at is not None
            ):
                raise _base_fixture_collision("analytics Opportunity")
        if len(target_markets) > 1:
            raise _base_fixture_collision("Target Market")
        if target_markets:
            target_market = target_markets[0]
            if owner_manifest_base is not None and owner_manifest_base.get("target_market_id") != target_market.id:
                raise _base_fixture_collision("Target Market identity")
            versions = list(
                (
                    await session.scalars(
                        select(ProspectTargetMarketVersion).where(
                            ProspectTargetMarketVersion.organisation_id == organisation_id,
                            ProspectTargetMarketVersion.target_market_id == target_market.id,
                        )
                    )
                ).all()
            )
            if (
                target_market.status != "active"
                or target_market.current_version != 1
                or target_market.created_by_user_id != owner_user_id
                or target_market.archived_at is not None
                or len(versions) != 1
            ):
                raise _base_fixture_collision("Target Market")
            version = versions[0]
            if (
                version.version != 1
                or version.description != "Synthetic large Australian organisations with distributed operations."
                or version.industries != ["Facilities services", "Healthcare", "Business software"]
                or version.countries != ["AU"]
                or version.regions != []
                or version.minimum_employee_band != "500_999"
                or version.organisation_types != []
                or version.preferred_business_characteristics != ["multi_site"]
                or version.excluded_industries != ["Retail"]
                or version.exclude_existing_accounts
                or version.research_objective != "Access-control and physical-security opportunity"
                or version.created_by_user_id != owner_user_id
            ):
                raise _base_fixture_collision("Target Market")
        timezone = organisation.timezone if organisation is not None else "UTC"
        expected_targets: dict[
            UUID,
            tuple[str, str, str, UUID | None, str, date, date, str | None, Decimal, UUID],
        ] = {
            target_ids["won-value"]: (
                "won_value",
                "personal",
                "self_set",
                owner_user_id,
                "year",
                date(2026, 1, 1),
                date(2026, 12, 31),
                "AUD",
                Decimal("1000000.00"),
                target_ids["won-value-revision"],
            ),
            target_ids["meetings"]: (
                "meetings_completed_count",
                "personal",
                "self_set",
                owner_user_id,
                "year",
                date(2026, 1, 1),
                date(2026, 12, 31),
                None,
                Decimal("12.00"),
                target_ids["meetings-revision"],
            ),
            target_ids["quarter-won-value"]: (
                "won_value",
                "organisation",
                "admin_assigned",
                None,
                "quarter",
                date(2026, 7, 1),
                date(2026, 9, 30),
                "AUD",
                Decimal("750000.00"),
                target_ids["quarter-won-value-revision"],
            ),
        }
        revisions_by_target = {item.target_id: item for item in sales_target_revisions}
        if len(sales_targets) != 3 or len(sales_target_revisions) != 3:
            raise _base_fixture_collision("Sales Target")
        for target in sales_targets:
            expected_target = expected_targets.get(target.id)
            if expected_target is None:
                raise _base_fixture_collision("Sales Target")
            metric, scope, origin, owner, period_type, period_start, period_end, currency, goal, revision_id = (
                expected_target
            )
            target_revision = revisions_by_target.get(target.id)
            if (
                target.metric_id != metric
                or target.metric_definition_version != "1"
                or target.scope != scope
                or target.origin != origin
                or target.owner_user_id != owner
                or target.pipeline_id is not None
                or target.period_type != period_type
                or target.period_start != period_start
                or target.period_end != period_end
                or target.timezone != timezone
                or target.currency != currency
                or target.created_by_user_id != owner_user_id
                or target.archived_at is not None
                or target_revision is None
                or target_revision.id != revision_id
                or target_revision.revision_number != 1
                or target_revision.goal_value != goal
                or target_revision.created_by_user_id != owner_user_id
            ):
                raise _base_fixture_collision("Sales Target")
        if (
            forecast_period is None
            or len(forecast_judgments) != 4
            or len(forecast_revisions) != 4
            or len(reviewer_judgments) != 3
            or len(reviewer_revisions) != 3
        ):
            raise _base_fixture_collision("Sales Forecast")
        if (
            forecast_period.period_type != "quarter"
            or forecast_period.period_start != date(2026, 7, 1)
            or forecast_period.period_end != date(2026, 9, 30)
            or forecast_period.timezone != timezone
            or forecast_period.created_by_user_id != owner_user_id
        ):
            raise _base_fixture_collision("Sales Forecast")
        expected_forecast: dict[str, tuple[UUID, str, str | None, int, int, str]] = {
            "main": (opportunity_id, "commit", "likely", 7, 5, "evaluation"),
            "likely": (analytics_opportunity_ids[14], "likely", "possible", 7, 5, "evaluation"),
            "possible": (analytics_opportunity_ids[16], "possible", "not_this_period", 7, 0, "negotiation"),
            "excluded": (analytics_opportunity_ids[19], "not_this_period", None, 7, 5, "evaluation"),
        }
        pipeline = await session.get(SalesPipeline, opportunity.pipeline_id)
        pipeline_stages = list(
            (
                await session.scalars(
                    select(SalesPipelineStage).where(
                        SalesPipelineStage.organisation_id == organisation_id,
                        SalesPipelineStage.pipeline_id == opportunity.pipeline_id,
                    )
                )
            ).all()
        )
        stage_by_key = {stage.stage_key: stage for stage in pipeline_stages}
        forecast_opportunities = {opportunity.id: opportunity, **analytics_by_id}
        if pipeline is None:
            raise _base_fixture_collision("Sales Forecast pipeline")
        judgments_by_id = {item.id: item for item in forecast_judgments}
        revisions_by_id = {item.id: item for item in forecast_revisions}
        reviewer_judgments_by_id = {item.id: item for item in reviewer_judgments}
        reviewer_revisions_by_id = {item.id: item for item in reviewer_revisions}
        for label, (
            selected_opportunity_id,
            category,
            reviewer_category,
            won_count,
            lost_count,
            stage_key,
        ) in expected_forecast.items():
            judgment = judgments_by_id.get(forecast_ids[f"{label}-judgment"])
            forecast_revision = revisions_by_id.get(forecast_ids[f"{label}-revision"])
            selected_opportunity = forecast_opportunities.get(selected_opportunity_id)
            expected_stage = stage_by_key.get(stage_key)
            if judgment is None or forecast_revision is None or selected_opportunity is None or expected_stage is None:
                raise _base_fixture_collision("Sales Forecast")
            if (
                judgment is not None
                and forecast_revision is not None
                and (
                    judgment.period_id != forecast_ids["period"]
                    or judgment.opportunity_id != selected_opportunity_id
                    or forecast_revision.judgment_id != judgment.id
                    or forecast_revision.revision_number != 1
                    or forecast_revision.category != category
                    or forecast_revision.created_by_user_id != owner_user_id
                    or forecast_revision.owner_user_id_snapshot != owner_user_id
                    or forecast_revision.amount_snapshot != selected_opportunity.estimated_value
                    or forecast_revision.currency_snapshot != selected_opportunity.currency
                    or forecast_revision.expected_close_date_snapshot != selected_opportunity.expected_close_date
                    or forecast_revision.pipeline_id_snapshot != pipeline.id
                    or forecast_revision.pipeline_name_snapshot != pipeline.name
                    or forecast_revision.stage_id_snapshot != expected_stage.id
                    or forecast_revision.stage_name_snapshot != expected_stage.name
                    or forecast_revision.opportunity_status_snapshot != "open"
                    or forecast_revision.model_version != "forecast_historical_stage_outcome_v1"
                    or forecast_revision.model_status
                    != ("available" if won_count + lost_count >= 10 else "insufficient_sample")
                    or forecast_revision.model_won_count != won_count
                    or forecast_revision.model_lost_count != lost_count
                    or forecast_revision.model_minimum_sample != 10
                    or forecast_revision.model_lookback_start != date(2024, 8, 25)
                    or forecast_revision.model_lookback_end != date(2026, 8, 24)
                )
            ):
                raise _base_fixture_collision("Sales Forecast")
            if reviewer_category is None:
                continue
            reviewer_judgment = reviewer_judgments_by_id.get(forecast_ids[f"{label}-reviewer-judgment"])
            reviewer_revision = reviewer_revisions_by_id.get(forecast_ids[f"{label}-reviewer-revision"])
            if reviewer_judgment is None or reviewer_revision is None:
                raise _base_fixture_collision("Sales Forecast reviewer")
            if (
                reviewer_judgment.period_id != forecast_ids["period"]
                or reviewer_judgment.opportunity_id != selected_opportunity_id
                or reviewer_revision.reviewer_judgment_id != reviewer_judgment.id
                or reviewer_revision.revision_number != 1
                or reviewer_revision.category != reviewer_category
                or reviewer_revision.created_by_user_id != owner_user_id
                or reviewer_revision.owner_user_id_snapshot != owner_user_id
                or reviewer_revision.amount_snapshot != selected_opportunity.estimated_value
                or reviewer_revision.currency_snapshot != selected_opportunity.currency
                or reviewer_revision.expected_close_date_snapshot != selected_opportunity.expected_close_date
                or reviewer_revision.pipeline_id_snapshot != pipeline.id
                or reviewer_revision.pipeline_name_snapshot != pipeline.name
                or reviewer_revision.stage_id_snapshot != expected_stage.id
                or reviewer_revision.stage_name_snapshot != expected_stage.name
                or reviewer_revision.opportunity_status_snapshot != "open"
                or reviewer_revision.model_version != "forecast_historical_stage_outcome_v1"
                or reviewer_revision.model_status
                != ("available" if won_count + lost_count >= 10 else "insufficient_sample")
                or reviewer_revision.model_won_count != won_count
                or reviewer_revision.model_lost_count != lost_count
                or reviewer_revision.model_minimum_sample != 10
                or reviewer_revision.model_lookback_start != date(2024, 8, 25)
                or reviewer_revision.model_lookback_end != date(2026, 8, 24)
            ):
                raise _base_fixture_collision("Sales Forecast reviewer")


async def _require_owner_demo_admin(
    session_factory: async_sessionmaker[AsyncSession],
    organisation_id: UUID,
    user_id: UUID,
) -> TenantContext:
    async with session_factory() as session:
        await set_tenant_database_context(session, organisation_id)
        membership = await session.get(OrganisationMembership, (organisation_id, user_id))
        if membership is None or membership.status != "active" or membership.role != "admin":
            raise ValueError("Owner demo refused: an active organisation administrator is required.")
        return TenantContext(
            organisation_id=organisation_id,
            user_id=user_id,
            role=cast(Role, membership.role),
        )


async def _owner_demo_reset_tenant(
    session_factory: async_sessionmaker[AsyncSession],
    organisation_id: UUID,
) -> TenantContext | None:
    company_id, _, _, _ = demo_ids(organisation_id)
    async with session_factory() as session:
        await set_tenant_database_context(session, organisation_id)
        company = await session.scalar(
            select(Company).where(Company.organisation_id == organisation_id, Company.id == company_id)
        )
        if company is None:
            return None
        membership = await session.get(OrganisationMembership, (organisation_id, company.owner_user_id))
        if membership is None or membership.status != "active" or membership.role != "admin":
            raise ValueError("Owner demo reset refused: the exact synthetic owner must remain an active administrator.")
        return TenantContext(
            organisation_id=organisation_id,
            user_id=company.owner_user_id,
            role=cast(Role, membership.role),
        )


def _base_manifest_encode(value: object) -> object:
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, dict):
        return {str(key): _base_manifest_encode(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_base_manifest_encode(item) for item in value]
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    raise ValueError("Owner demo refused an unsupported base-manifest value.")


def _base_manifest_decode(value: object) -> object:
    if isinstance(value, dict):
        return {str(key): _base_manifest_decode(item) for key, item in value.items()}
    if isinstance(value, list):
        return tuple(_base_manifest_decode(item) for item in value)
    if isinstance(value, str):
        try:
            return UUID(value)
        except ValueError:
            return value
    return value


def _validated_base_manifest(
    event: BetaSystemEvent,
    organisation_id: UUID,
    user_id: UUID,
) -> dict[str, object]:
    expected_id = owner_demo_base_manifest_id(organisation_id)
    if (
        event.id != expected_id
        or event.organisation_id != organisation_id
        or event.actor_user_id != user_id
        or event.event_type != OWNER_BASE_MANIFEST_EVENT
        or event.subject_id != expected_id
        or set(event.metadata_json) != {"schemaVersion", "base", "extensions"}
        or event.metadata_json.get("schemaVersion") != OWNER_BASE_MANIFEST_SCHEMA_VERSION
        or not isinstance(event.metadata_json.get("base"), dict)
    ):
        raise _base_fixture_collision("owner manifest")
    decoded = _base_manifest_decode(event.metadata_json["base"])
    if not isinstance(decoded, dict):
        raise _base_fixture_collision("owner manifest")
    company_id, opportunity_id, _, _ = demo_ids(organisation_id)
    if (
        decoded.get("status") != "ready"
        or decoded.get("company_id") != company_id
        or decoded.get("opportunity_id") != opportunity_id
        or decoded.get("contact_id") != demo_phone_contact_id(organisation_id)
        or decoded.get("campaign_id") != demo_campaign_ids(organisation_id)["campaign"]
        or decoded.get("provider_calls") != 0
    ):
        raise _base_fixture_collision("owner manifest")
    return decoded


def _validated_base_reservation(
    event: BetaSystemEvent,
    organisation_id: UUID,
    user_id: UUID,
) -> _MethodologyReservation:
    expected_id = owner_demo_base_manifest_id(organisation_id)
    if (
        event.id != expected_id
        or event.organisation_id != organisation_id
        or event.actor_user_id != user_id
        or event.event_type != OWNER_BASE_RESERVATION_EVENT
        or event.subject_id != expected_id
        or set(event.metadata_json)
        != {
            "schemaVersion",
            "methodologySettingPreexisting",
            "methodologySettingOriginal",
        }
        or event.metadata_json.get("schemaVersion") != OWNER_BASE_MANIFEST_SCHEMA_VERSION
        or not isinstance(event.metadata_json.get("methodologySettingPreexisting"), bool)
    ):
        raise _base_fixture_collision("owner reservation")
    preexisting = cast(bool, event.metadata_json["methodologySettingPreexisting"])
    raw_original = event.metadata_json.get("methodologySettingOriginal")
    if not preexisting:
        if raw_original is not None:
            raise _base_fixture_collision("owner reservation")
        return _MethodologyReservation(preexisting=False, original=None)
    if not isinstance(raw_original, dict) or set(raw_original) != {
        "selection",
        "customDefinitionId",
        "updatedByUserId",
        "createdAt",
        "updatedAt",
    }:
        raise _base_fixture_collision("owner reservation")
    selection = raw_original.get("selection")
    custom_definition_id_raw = raw_original.get("customDefinitionId")
    updated_by_user_id_raw = raw_original.get("updatedByUserId")
    created_at_raw = raw_original.get("createdAt")
    updated_at_raw = raw_original.get("updatedAt")
    if (
        selection not in {"none", "meddic", "meddicc", "meddpicc", "bant", "spiced", "custom"}
        or (selection == "custom") != (custom_definition_id_raw is not None)
        or (custom_definition_id_raw is not None and not isinstance(custom_definition_id_raw, str))
        or not isinstance(updated_by_user_id_raw, str)
        or not isinstance(created_at_raw, str)
        or not isinstance(updated_at_raw, str)
    ):
        raise _base_fixture_collision("owner reservation")
    try:
        custom_definition_id = UUID(custom_definition_id_raw) if isinstance(custom_definition_id_raw, str) else None
        updated_by_user_id = UUID(updated_by_user_id_raw)
        created_at = datetime.fromisoformat(created_at_raw)
        updated_at = datetime.fromisoformat(updated_at_raw)
    except ValueError as exc:
        raise _base_fixture_collision("owner reservation") from exc
    return _MethodologyReservation(
        preexisting=True,
        original=_MethodologySettingSnapshot(
            selection=cast(str, selection),
            custom_definition_id=custom_definition_id,
            updated_by_user_id=updated_by_user_id,
            created_at=created_at,
            updated_at=updated_at,
        ),
    )


def _methodology_setting_snapshot(
    setting: OrganisationMethodologySetting,
) -> _MethodologySettingSnapshot:
    return _MethodologySettingSnapshot(
        selection=setting.selection,
        custom_definition_id=setting.custom_definition_id,
        updated_by_user_id=setting.updated_by_user_id,
        created_at=setting.created_at,
        updated_at=setting.updated_at,
    )


def _methodology_reservation_metadata(
    reservation: _MethodologyReservation,
) -> dict[str, object]:
    original = reservation.original
    return {
        "schemaVersion": OWNER_BASE_MANIFEST_SCHEMA_VERSION,
        "methodologySettingPreexisting": reservation.preexisting,
        "methodologySettingOriginal": (
            {
                "selection": original.selection,
                "customDefinitionId": (
                    str(original.custom_definition_id) if original.custom_definition_id is not None else None
                ),
                "updatedByUserId": str(original.updated_by_user_id),
                "createdAt": original.created_at.isoformat(),
                "updatedAt": original.updated_at.isoformat(),
            }
            if original is not None
            else None
        ),
    }


def _methodology_setting_matches_snapshot(
    setting: OrganisationMethodologySetting,
    snapshot: _MethodologySettingSnapshot,
) -> bool:
    return (
        setting.selection == snapshot.selection
        and setting.custom_definition_id == snapshot.custom_definition_id
        and setting.updated_by_user_id == snapshot.updated_by_user_id
        and setting.created_at == snapshot.created_at
        and setting.updated_at == snapshot.updated_at
    )


def _assert_recoverable_methodology_reservation_state(
    setting: OrganisationMethodologySetting | None,
    user_id: UUID,
    reservation: _MethodologyReservation,
) -> None:
    if setting is None:
        raise _base_fixture_collision("owner methodology setting reservation")
    if reservation.original is not None and _methodology_setting_matches_snapshot(
        setting,
        reservation.original,
    ):
        return
    if (
        setting.selection not in {"bant", "meddpicc"}
        or setting.custom_definition_id is not None
        or setting.updated_by_user_id != user_id
    ):
        raise _base_fixture_collision("owner methodology setting reservation")


async def _restore_methodology_reservation_state(
    session: AsyncSession,
    setting: OrganisationMethodologySetting | None,
    user_id: UUID,
    reservation: _MethodologyReservation,
    *,
    remove_owner_created: bool,
) -> None:
    _assert_recoverable_methodology_reservation_state(setting, user_id, reservation)
    assert setting is not None
    original = reservation.original
    if original is None:
        if remove_owner_created:
            await session.delete(setting)
            return
        setting.selection = "meddpicc"
        setting.custom_definition_id = None
        setting.updated_by_user_id = user_id
        return
    setting.selection = original.selection
    setting.custom_definition_id = original.custom_definition_id
    setting.updated_by_user_id = original.updated_by_user_id
    setting.created_at = original.created_at
    setting.updated_at = original.updated_at


def _validate_completed_owner_methodology_setting(
    setting: OrganisationMethodologySetting | None,
    user_id: UUID,
    base: dict[str, object],
    *,
    allow_missing: bool,
) -> None:
    preexisting = base.get("methodology_setting_preexisting")
    if preexisting is None or preexisting is True:
        return
    if preexisting is not False:
        raise _base_fixture_collision("owner methodology setting manifest")
    if setting is None and allow_missing:
        return
    if setting is None or (
        setting.selection != "meddpicc"
        or setting.custom_definition_id is not None
        or setting.updated_by_user_id != user_id
    ):
        raise _base_fixture_collision("owner methodology setting manifest")


def _validated_extension_manifest(
    event: BetaSystemEvent,
    organisation_id: UUID,
    user_id: UUID,
) -> dict[str, UUID] | None:
    _validated_base_manifest(event, organisation_id, user_id)
    raw_extensions = event.metadata_json.get("extensions")
    if raw_extensions is None:
        return None
    if not isinstance(raw_extensions, dict) or set(raw_extensions) != OWNER_EXTENSION_MANIFEST_KEYS:
        raise _base_fixture_collision("owner extension manifest")
    extensions: dict[str, UUID] = {}
    for key, raw_value in raw_extensions.items():
        try:
            extensions[key] = UUID(str(raw_value))
        except ValueError as exc:
            raise _base_fixture_collision("owner extension manifest") from exc
    return extensions


async def _reset_owner_methodology_setting(
    session: AsyncSession,
    event: BetaSystemEvent,
    organisation_id: UUID,
) -> None:
    actor_user_id = event.actor_user_id
    if actor_user_id is None:
        raise _base_fixture_collision("owner manifest")
    setting = await session.get(OrganisationMethodologySetting, organisation_id)
    if event.event_type == OWNER_BASE_RESERVATION_EVENT:
        reservation = _validated_base_reservation(event, organisation_id, actor_user_id)
        await _restore_methodology_reservation_state(
            session,
            setting,
            actor_user_id,
            reservation,
            remove_owner_created=True,
        )
        return
    base = _validated_base_manifest(event, organisation_id, actor_user_id)
    _validate_completed_owner_methodology_setting(
        setting,
        actor_user_id,
        base,
        allow_missing=True,
    )
    if base.get("methodology_setting_preexisting") is not False or setting is None:
        return None
    await session.delete(setting)


async def _owner_base_fixture(
    session_factory: async_sessionmaker[AsyncSession],
    tenant: TenantContext,
    settings: Settings,
) -> dict[str, object]:
    manifest_id = owner_demo_base_manifest_id(tenant.organisation_id)
    async with session_factory() as session, session.begin():
        await set_tenant_database_context(session, tenant.organisation_id)
        methodology_setting = await session.get(
            OrganisationMethodologySetting,
            tenant.organisation_id,
        )
        existing = await session.scalar(
            select(BetaSystemEvent)
            .where(
                BetaSystemEvent.organisation_id == tenant.organisation_id,
                BetaSystemEvent.id == manifest_id,
            )
            .with_for_update()
        )
        if existing is not None:
            if existing.event_type == OWNER_BASE_MANIFEST_EVENT:
                base = _validated_base_manifest(existing, tenant.organisation_id, tenant.user_id)
                _validate_completed_owner_methodology_setting(
                    methodology_setting,
                    tenant.user_id,
                    base,
                    allow_missing=False,
                )
                return base
            methodology_reservation = _validated_base_reservation(
                existing,
                tenant.organisation_id,
                tenant.user_id,
            )
            await _restore_methodology_reservation_state(
                session,
                methodology_setting,
                tenant.user_id,
                methodology_reservation,
                remove_owner_created=False,
            )
        else:
            original_methodology_setting = (
                _methodology_setting_snapshot(methodology_setting) if methodology_setting is not None else None
            )
            methodology_reservation = _MethodologyReservation(
                preexisting=original_methodology_setting is not None,
                original=original_methodology_setting,
            )
            if methodology_setting is None:
                # Establish ownership atomically with the reservation.  A crash
                # can therefore never leave a newly created global setting
                # without the durable proof needed to remove it safely.
                session.add(
                    OrganisationMethodologySetting(
                        organisation_id=tenant.organisation_id,
                        selection="meddpicc",
                        updated_by_user_id=tenant.user_id,
                    )
                )
            session.add(
                BetaSystemEvent(
                    id=manifest_id,
                    organisation_id=tenant.organisation_id,
                    actor_user_id=tenant.user_id,
                    event_type=OWNER_BASE_RESERVATION_EVENT,
                    subject_id=manifest_id,
                    metadata_json=_methodology_reservation_metadata(methodology_reservation),
                )
            )
    seeded = await seed_demo_data(
        session_factory,
        tenant.organisation_id,
        tenant.user_id,
        settings,
        preserve_existing_methodology_setting=True,
    )
    seeded["methodology_setting_preexisting"] = methodology_reservation.preexisting
    encoded = _base_manifest_encode(seeded)
    if not isinstance(encoded, dict):
        raise ValueError("Owner demo refused an invalid base-manifest value.")
    async with session_factory() as session, session.begin():
        await set_tenant_database_context(session, tenant.organisation_id)
        existing = await session.scalar(
            select(BetaSystemEvent)
            .where(
                BetaSystemEvent.organisation_id == tenant.organisation_id,
                BetaSystemEvent.id == manifest_id,
            )
            .with_for_update()
        )
        if existing is None:
            raise _base_fixture_collision("owner reservation")
        if existing.event_type == OWNER_BASE_MANIFEST_EVENT:
            base = _validated_base_manifest(existing, tenant.organisation_id, tenant.user_id)
            methodology_setting = await session.get(
                OrganisationMethodologySetting,
                tenant.organisation_id,
            )
            _validate_completed_owner_methodology_setting(
                methodology_setting,
                tenant.user_id,
                base,
                allow_missing=False,
            )
            return base
        reserved_methodology = _validated_base_reservation(
            existing,
            tenant.organisation_id,
            tenant.user_id,
        )
        if reserved_methodology != methodology_reservation:
            raise _base_fixture_collision("owner reservation")
        methodology_setting = await session.get(
            OrganisationMethodologySetting,
            tenant.organisation_id,
        )
        if methodology_reservation.original is not None:
            if methodology_setting is None or not _methodology_setting_matches_snapshot(
                methodology_setting,
                methodology_reservation.original,
            ):
                raise _base_fixture_collision("owner methodology setting reservation")
        elif methodology_setting is None or (
            methodology_setting.selection != "meddpicc"
            or methodology_setting.custom_definition_id is not None
            or methodology_setting.updated_by_user_id != tenant.user_id
        ):
            raise _base_fixture_collision("owner methodology setting reservation")
        existing.event_type = OWNER_BASE_MANIFEST_EVENT
        existing.metadata_json = {
            "schemaVersion": OWNER_BASE_MANIFEST_SCHEMA_VERSION,
            "base": encoded,
            "extensions": None,
        }
    decoded = _base_manifest_decode(encoded)
    if not isinstance(decoded, dict):
        raise ValueError("Owner demo refused an invalid base-manifest value.")
    return decoded


async def _finalise_owner_manifest(
    session_factory: async_sessionmaker[AsyncSession],
    tenant: TenantContext,
    extensions: dict[str, UUID],
) -> None:
    if set(extensions) != OWNER_EXTENSION_MANIFEST_KEYS:
        raise ValueError("Owner demo refused an incomplete extension manifest.")
    manifest_id = owner_demo_base_manifest_id(tenant.organisation_id)
    encoded_extensions = {key: str(value) for key, value in extensions.items()}
    async with session_factory() as session, session.begin():
        await set_tenant_database_context(session, tenant.organisation_id)
        event = await session.scalar(
            select(BetaSystemEvent)
            .where(
                BetaSystemEvent.organisation_id == tenant.organisation_id,
                BetaSystemEvent.id == manifest_id,
            )
            .with_for_update()
        )
        if event is None:
            raise _base_fixture_collision("owner manifest")
        existing = _validated_extension_manifest(event, tenant.organisation_id, tenant.user_id)
        if existing is not None and existing != extensions:
            raise _base_fixture_collision("owner extension manifest")
        event.metadata_json = {
            "schemaVersion": OWNER_BASE_MANIFEST_SCHEMA_VERSION,
            "base": event.metadata_json["base"],
            "extensions": encoded_extensions,
        }


def _selling_profile_content() -> SellingProfileContent:
    return SellingProfileContent.model_validate(
        {
            "companyDescription": (
                "Oryntela helps GTM teams prepare for customer conversations, preserve trustworthy notes and "
                "follow through with evidence-backed context across long relationships."
            ),
            "offerings": [
                {
                    "name": "Oryntela Conversation Intelligence",
                    "description": (
                        "A calm, source-aware workspace for preparation, meetings, phone calls, notes, "
                        "framework review and seller-approved follow-up."
                    ),
                    "whoNormallyBuys": ["Founders", "Revenue leaders", "Relationship-led sales teams"],
                    "problemsSolved": [
                        "Scattered conversation context",
                        "Important commitments lost between meetings",
                        "Generic preparation that ignores prior conversations",
                    ],
                    "intendedOutcomes": [
                        "Better prepared customer conversations",
                        "Personal notes that remain intact",
                        "Clearer reviewed follow-up and unresolved questions",
                    ],
                    "differentiators": [
                        "Evidence-aware human review",
                        "Source-labelled conversation memory",
                        "Meetings, calls and in-person notes in one restrained workflow",
                    ],
                    "competitorsAlternatives": ["Manual notes", "CRM-only workflows", "Disconnected AI tools"],
                    "commonObjections": [
                        "We already take notes — Oryntela keeps the original notes separate and makes prior context usable.",
                        "We do not want another CRM — Opportunities are optional in the conversation workflow.",
                    ],
                    "approvedProof": ["Approved synthetic owner demonstration"],
                    "approvedClaims": [
                        "Keeps approved conversation context available to authorised organisation members",
                        "Preserves personal notes when later processing is unavailable",
                    ],
                }
            ],
        }
    )


OWNER_KNOWLEDGE_URL = "https://northstar-facilities.example/"
OWNER_KNOWLEDGE_TEXT = (
    "Synthetic approved context: Oryntela supports preparation, personal notes, evidence review and follow-up "
    "for meetings, phone calls and in-person conversations. It does not replace a CRM, fabricate buyer facts "
    "or send follow-up without an authorised user action."
)
OWNER_KNOWLEDGE_PROCESSING_STATES = frozenset(
    {"queued", "fetching", "extracting", "ready_for_review", "approved", "stale", "failed"}
)


async def _seed_company_knowledge_sources(
    service: SellingProfileService,
) -> tuple[UUID, UUID]:
    requests = (
        CompanyKnowledgeSourceCreate(
            source_kind="website",
            label="[DEMO] Oryntela website address",
            website_url=OWNER_KNOWLEDGE_URL,
            source_text=None,
            authority_confirmed=True,
            idempotency_key=OWNER_KNOWLEDGE_WEBSITE_KEY,
        ),
        CompanyKnowledgeSourceCreate(
            source_kind="text_file",
            label="[DEMO] Approved conversation product brief.txt",
            website_url=None,
            source_text=OWNER_KNOWLEDGE_TEXT,
            authority_confirmed=True,
            idempotency_key=OWNER_KNOWLEDGE_TEXT_KEY,
        ),
    )
    saved = (
        await service.create_knowledge_source(requests[0]),
        await service.create_knowledge_source(requests[1]),
    )
    expected = (
        (
            "website",
            "[DEMO] Oryntela website address",
            OWNER_KNOWLEDGE_URL,
            0,
            hashlib.sha256(OWNER_KNOWLEDGE_URL.encode()).hexdigest(),
        ),
        (
            "text_file",
            "[DEMO] Approved conversation product brief.txt",
            None,
            len(OWNER_KNOWLEDGE_TEXT),
            hashlib.sha256(OWNER_KNOWLEDGE_TEXT.encode()).hexdigest(),
        ),
    )
    for source, values in zip(saved, expected, strict=True):
        if (
            source.source_kind,
            source.label,
            source.website_url,
            source.character_count,
            source.content_fingerprint,
        ) != values or source.processing_status not in OWNER_KNOWLEDGE_PROCESSING_STATES:
            raise _owner_collision("Company Knowledge source")
    return saved[0].id, saved[1].id


def _value_model_definition() -> ValueModelDefinition:
    return ValueModelDefinition.model_validate(
        {
            "inputs": [
                {
                    "key": "users",
                    "label": "Users supported",
                    "description": "Synthetic number of users supported by the proposed workflow.",
                    "valueType": "count",
                    "unit": "count",
                    "minimum": "1",
                    "maximum": "100000",
                    "decimalPrecision": 0,
                    "sourcePolicy": "reviewed_manual",
                    "material": True,
                    "sensitivityEligible": False,
                    "displayOrder": 1,
                },
                {
                    "key": "hours_saved_per_user_month",
                    "label": "Hours saved per user each month",
                    "description": "Synthetic seller-reviewed time assumption for this demonstration.",
                    "valueType": "hours",
                    "unit": "hours",
                    "minimum": "0",
                    "maximum": "100",
                    "decimalPrecision": 1,
                    "sourcePolicy": "reviewed_manual",
                    "material": True,
                    "sensitivityEligible": True,
                    "displayOrder": 2,
                },
                {
                    "key": "labour_cost_per_hour",
                    "label": "Loaded labour cost",
                    "description": "Approved synthetic organisation assumption for the demonstration.",
                    "valueType": "currency",
                    "unit": "currency_per_hour",
                    "minimum": "0",
                    "maximum": "10000",
                    "decimalPrecision": 2,
                    "defaultValue": "65",
                    "defaultOrigin": "organisation_assumption",
                    "defaultSourceReference": "Approved synthetic owner-demo assumption.",
                    "assumptionLocked": True,
                    "sourcePolicy": "approved_org_only",
                    "material": True,
                    "sensitivityEligible": False,
                    "displayOrder": 3,
                },
                {
                    "key": "annual_subscription_cost",
                    "label": "Annual subscription cost",
                    "description": "Illustrative synthetic commercial assumption, not a live quote.",
                    "valueType": "currency",
                    "unit": "currency_per_year",
                    "minimum": "0",
                    "maximum": "10000000",
                    "decimalPrecision": 2,
                    "defaultValue": "48000",
                    "defaultOrigin": "organisation_assumption",
                    "defaultSourceReference": "Illustrative synthetic owner-demo assumption.",
                    "assumptionLocked": True,
                    "sourcePolicy": "approved_org_only",
                    "material": True,
                    "sensitivityEligible": False,
                    "displayOrder": 4,
                },
                {
                    "key": "implementation_cost",
                    "label": "Implementation cost",
                    "description": "Illustrative synthetic implementation assumption, not a live quote.",
                    "valueType": "currency",
                    "unit": "currency",
                    "minimum": "0",
                    "maximum": "10000000",
                    "decimalPrecision": 2,
                    "defaultValue": "20000",
                    "defaultOrigin": "organisation_assumption",
                    "defaultSourceReference": "Illustrative synthetic owner-demo assumption.",
                    "assumptionLocked": True,
                    "sourcePolicy": "approved_org_only",
                    "material": True,
                    "sensitivityEligible": False,
                    "displayOrder": 5,
                },
            ],
            "outputs": [
                {
                    "key": "annual_hours_saved",
                    "label": "Annual hours saved",
                    "description": "Modelled annual capacity released from the supplied assumptions.",
                    "formula": "users * hours_saved_per_user_month * 12",
                    "unit": "hours_per_year",
                    "displayPrecision": 1,
                    "highlight": True,
                    "displayOrder": 1,
                },
                {
                    "key": "annual_labour_value",
                    "label": "Annual labour value",
                    "description": "Modelled annual value of the released capacity.",
                    "formula": "annual_hours_saved * labour_cost_per_hour",
                    "unit": "currency_per_year",
                    "displayPrecision": 2,
                    "highlight": True,
                    "displayOrder": 2,
                },
                {
                    "key": "first_year_cost",
                    "label": "First-year cost",
                    "description": "Illustrative subscription and implementation assumptions combined.",
                    "formula": "annual_subscription_cost + implementation_cost",
                    "unit": "currency",
                    "displayPrecision": 2,
                    "highlight": False,
                    "scenarioSensitive": False,
                    "displayOrder": 3,
                },
                {
                    "key": "first_year_net_benefit",
                    "label": "First-year net benefit",
                    "description": "Modelled annual labour value less illustrative first-year cost.",
                    "formula": "annual_labour_value - first_year_cost",
                    "unit": "currency",
                    "displayPrecision": 2,
                    "highlight": True,
                    "displayOrder": 4,
                },
                {
                    "key": "roi_percentage",
                    "label": "First-year ROI",
                    "description": "Modelled first-year net benefit divided by first-year cost.",
                    "formula": "safe_divide(first_year_net_benefit, first_year_cost) * 100",
                    "unit": "percentage",
                    "displayPrecision": 1,
                    "highlight": True,
                    "displayOrder": 5,
                },
            ],
            "customerDisclaimer": (
                "Synthetic demonstration based on the visible inputs and assumptions shown; not a quote or a "
                "guarantee of future results."
            ),
        }
    )


def _business_case_inputs() -> list[BusinessCaseInputValue]:
    return [
        BusinessCaseInputValue(key="users", value="75", origin="salesperson_reported"),
        BusinessCaseInputValue(
            key="hours_saved_per_user_month",
            value="3",
            origin="salesperson_reported",
        ),
        BusinessCaseInputValue(
            key="labour_cost_per_hour",
            value="65",
            origin="organisation_assumption",
        ),
        BusinessCaseInputValue(
            key="annual_subscription_cost",
            value="48000",
            origin="organisation_assumption",
        ),
        BusinessCaseInputValue(
            key="implementation_cost",
            value="20000",
            origin="organisation_assumption",
        ),
    ]


def _business_case_calculation_request() -> BusinessCaseCalculateRequest:
    return BusinessCaseCalculateRequest(
        inputs=_business_case_inputs(),
        scenarios=[
            BusinessCaseScenarioRequest(
                name="conservative",
                overrides=[ScenarioOverride(key="hours_saved_per_user_month", value="2")],
            ),
            BusinessCaseScenarioRequest(
                name="upside",
                overrides=[ScenarioOverride(key="hours_saved_per_user_month", value="4")],
            ),
        ],
        sensitivity=SensitivityRequest(
            input_key="hours_saved_per_user_month",
            values=["2", "3", "4"],
        ),
        idempotency_key=OWNER_BUSINESS_CASE_CALCULATION_KEY,
    )


def _owner_collision(resource: str) -> ValueError:
    return ValueError(f"Owner demo refused: {resource} identity/content collision.")


def _json_fingerprint(value: object) -> str:
    encoded = json.dumps(value, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


async def _validated_owner_profile(
    session: AsyncSession,
    tenant: TenantContext,
) -> SellingProfileRevision | None:
    owner_revisions = list(
        (
            await session.scalars(
                select(SellingProfileRevision).where(
                    SellingProfileRevision.organisation_id == tenant.organisation_id,
                    SellingProfileRevision.idempotency_key == OWNER_PROFILE_KEY,
                )
            )
        ).all()
    )
    all_revisions = list(
        (
            await session.scalars(
                select(SellingProfileRevision).where(SellingProfileRevision.organisation_id == tenant.organisation_id)
            )
        ).all()
    )
    profile = await session.scalar(
        select(SellingProfile).where(SellingProfile.organisation_id == tenant.organisation_id)
    )
    if not owner_revisions:
        if profile is not None or all_revisions:
            raise _owner_collision("Company & Selling Profile")
        return None
    if len(owner_revisions) != 1 or len(all_revisions) != 1:
        raise _owner_collision("Company & Selling Profile")
    revision = owner_revisions[0]
    expected_content = _selling_profile_content().model_dump(mode="json", by_alias=True)
    expected_fingerprint = _json_fingerprint(expected_content)
    if (
        profile is None
        or revision.profile_id != profile.id
        or profile.created_by_user_id != tenant.user_id
        or revision.created_by_user_id != tenant.user_id
        or revision.revision_number != 1
        or revision.schema_version != 1
        or revision.lock_version != 1
        or revision.content_json != expected_content
        or revision.content_fingerprint != expected_fingerprint
        or revision.state not in {"draft", "approved"}
        or (revision.state == "draft" and revision.approved_by_user_id is not None)
        or (revision.state == "approved" and revision.approved_by_user_id != tenant.user_id)
    ):
        raise _owner_collision("Company & Selling Profile")
    return revision


async def _validated_owner_value_model(
    session: AsyncSession,
    tenant: TenantContext,
    service: BusinessCaseService,
) -> tuple[CreateValueModel, CreateValueModelVersion] | None:
    models = list(
        (
            await session.scalars(
                select(CreateValueModel).where(
                    CreateValueModel.organisation_id == tenant.organisation_id,
                    CreateValueModel.idempotency_key == OWNER_VALUE_MODEL_KEY,
                )
            )
        ).all()
    )
    if not models:
        return None
    if len(models) != 1:
        raise _owner_collision("Value Model")
    model = models[0]
    versions = list(
        (
            await session.scalars(
                select(CreateValueModelVersion).where(
                    CreateValueModelVersion.organisation_id == tenant.organisation_id,
                    CreateValueModelVersion.model_id == model.id,
                )
            )
        ).all()
    )
    validated, expected_definition, expected_fingerprint = service._validated_definition(_value_model_definition())
    if (
        model.name != OWNER_VALUE_MODEL_NAME
        or model.description != OWNER_VALUE_MODEL_DESCRIPTION
        or model.state != "active"
        or model.created_by_user_id != tenant.user_id
        or len(versions) != 1
    ):
        raise _owner_collision("Value Model")
    version = versions[0]
    if (
        version.version != 1
        or version.idempotency_key != OWNER_VALUE_MODEL_KEY
        or version.created_by_user_id != tenant.user_id
        or version.definition_json != expected_definition
        or version.canonical_ast_json != validated.canonical_ast
        or version.formula_engine_version != "bounded_decimal_v1"
        or version.fingerprint != expected_fingerprint
        or version.state not in {"draft", "approved"}
        or (version.state == "draft" and version.approved_by_user_id is not None)
        or (version.state == "approved" and version.approved_by_user_id != tenant.user_id)
    ):
        raise _owner_collision("Value Model")
    return model, version


def _normalised_business_case_inputs(
    service: BusinessCaseService,
    definition: ValueModelDefinition,
) -> list[dict[str, object]]:
    requested = {item.key: item for item in _business_case_inputs()}
    expected: list[dict[str, object]] = []
    for item in sorted(definition.inputs, key=lambda value: value.display_order):
        supplied = requested[item.key]
        value = Decimal(supplied.value)
        expected.append(
            {
                "key": item.key,
                "label": item.label,
                "value": service._decimal_string(value),
                "calculationValue": service._decimal_string(
                    value / Decimal("100") if item.unit == "percentage" else value
                ),
                "unit": item.unit,
                "origin": supplied.origin,
                "sourceId": None,
                "sourceLabel": (
                    "Approved organisation assumption"
                    if supplied.origin == "organisation_assumption"
                    else "Reported by you"
                ),
                "assumption": supplied.origin in {"organisation_assumption", "user_entered", "unknown"},
                "material": item.material,
                "customerFacing": item.customer_facing,
                "freshness": "current",
            }
        )
    return expected


def _normalise_stored_business_case_inputs(values: list[object]) -> list[dict[str, object]]:
    normalised: list[dict[str, object]] = []
    for raw in values:
        if not isinstance(raw, dict):
            raise _owner_collision("Business Case")
        item = dict(raw)
        observed_at = item.pop("observedAt", None)
        if not isinstance(observed_at, str):
            raise _owner_collision("Business Case")
        try:
            parsed = datetime.fromisoformat(observed_at.replace("Z", "+00:00"))
        except ValueError as exc:
            raise _owner_collision("Business Case") from exc
        if parsed.utcoffset() is None:
            raise _owner_collision("Business Case")
        normalised.append(item)
    return normalised


async def _validated_owner_business_case(
    session: AsyncSession,
    tenant: TenantContext,
    service: BusinessCaseService,
    company_id: UUID,
    opportunity_id: UUID,
    model: CreateValueModel,
    model_version: CreateValueModelVersion,
) -> tuple[CreateBusinessCase, CreateBusinessCaseVersion | None] | None:
    cases = list(
        (
            await session.scalars(
                select(CreateBusinessCase).where(
                    CreateBusinessCase.organisation_id == tenant.organisation_id,
                    CreateBusinessCase.idempotency_key == OWNER_BUSINESS_CASE_KEY,
                )
            )
        ).all()
    )
    if not cases:
        return None
    if len(cases) != 1:
        raise _owner_collision("Business Case")
    business_case = cases[0]
    versions = list(
        (
            await session.scalars(
                select(CreateBusinessCaseVersion).where(
                    CreateBusinessCaseVersion.organisation_id == tenant.organisation_id,
                    CreateBusinessCaseVersion.case_id == business_case.id,
                )
            )
        ).all()
    )
    if (
        business_case.account_id != company_id
        or business_case.opportunity_id != opportunity_id
        or business_case.model_id != model.id
        or business_case.model_version_id != model_version.id
        or business_case.created_by_user_id != tenant.user_id
        or business_case.title != OWNER_BUSINESS_CASE_TITLE
        or business_case.currency != "AUD"
        or len(versions) > 1
    ):
        raise _owner_collision("Business Case")
    if not versions:
        if business_case.state != "draft":
            raise _owner_collision("Business Case")
        return business_case, None

    version = versions[0]
    definition = _value_model_definition()
    validated, _, expected_model_fingerprint = service._validated_definition(definition)
    request = _business_case_calculation_request()
    values = {item.key: Decimal(item.value) for item in request.inputs}
    scenarios = [
        ScenarioCalculationResponse(
            name="base",
            overrides=[],
            outputs=service._calculate_outputs(validated, definition, values),
        )
    ]
    scenario_fingerprint: list[dict[str, object]] = []
    for requested_scenario in request.scenarios:
        scenario_values = dict(values)
        for override in requested_scenario.overrides:
            scenario_values[override.key] = Decimal(override.value)
        scenarios.append(
            ScenarioCalculationResponse(
                name=requested_scenario.name,
                overrides=requested_scenario.overrides,
                outputs=service._calculate_outputs(validated, definition, scenario_values),
            )
        )
        scenario_fingerprint.append(
            {
                "name": requested_scenario.name,
                "overrides": [item.model_dump(mode="json") for item in requested_scenario.overrides],
            }
        )
    sensitivity = service._calculate_sensitivity(validated, definition, values, request)
    fingerprint_inputs = [
        {
            "key": item.key,
            "value": service._decimal_string(Decimal(item.value)),
            "origin": item.origin,
            "sourceId": None,
            "observedAt": None,
        }
        for item in sorted(
            request.inputs,
            key=lambda value: next(
                definition_item.display_order
                for definition_item in definition.inputs
                if definition_item.key == value.key
            ),
        )
    ]
    expected_calculation_fingerprint = service._hash_json(
        {
            "modelVersionId": str(model_version.id),
            "modelFingerprint": model_version.fingerprint,
            "currency": "AUD",
            "inputs": fingerprint_inputs,
            "scenarios": scenario_fingerprint,
            "sensitivity": request.sensitivity.model_dump(mode="json") if request.sensitivity else None,
        }
    )
    expected_scenarios = [item.model_dump(mode="json", by_alias=True) for item in scenarios]
    expected_sensitivity = sensitivity.model_dump(mode="json", by_alias=True) if sensitivity else None
    expected_lineage = service._lineage(definition, model_version, scenarios)
    valid_state = (
        business_case.state == "calculated"
        and version.review_state == "pending"
        and version.approved_by_user_id is None
    ) or (
        business_case.state == "approved"
        and version.review_state == "approved"
        and version.approved_by_user_id == tenant.user_id
    )
    if (
        version.version != 1
        or version.model_id != model.id
        or version.model_version_id != model_version.id
        or version.currency != "AUD"
        or version.formula_engine_version != "bounded_decimal_v1"
        or version.model_fingerprint != expected_model_fingerprint
        or version.calculation_fingerprint != expected_calculation_fingerprint
        or _normalise_stored_business_case_inputs(version.inputs_json)
        != _normalised_business_case_inputs(service, definition)
        or version.scenarios_json != expected_scenarios
        or version.sensitivity_json != expected_sensitivity
        or version.lineage_json != expected_lineage
        or version.created_by_user_id != tenant.user_id
        or version.idempotency_key != OWNER_BUSINESS_CASE_CALCULATION_KEY
        or not valid_state
    ):
        raise _owner_collision("Business Case")
    return business_case, version


def _normalise_plan(plan: Sequence[object]) -> list[dict[str, object]]:
    normalised: list[dict[str, object]] = []
    for raw in plan:
        if not isinstance(raw, dict):
            raise _owner_collision("Presentation")
        item = dict(raw)
        try:
            UUID(str(item.pop("id")))
        except (KeyError, ValueError) as exc:
            raise _owner_collision("Presentation") from exc
        normalised.append(item)
    return normalised


def _claims_match_deterministic_output(
    actual_claims: list[object],
    expected_claims: list[dict[str, object]],
    *,
    approved: bool,
) -> bool:
    if len(actual_claims) != len(expected_claims):
        return False
    for raw_actual, raw_expected in zip(actual_claims, expected_claims, strict=True):
        if not isinstance(raw_actual, dict):
            return False
        actual = dict(raw_actual)
        expected = dict(raw_expected)
        try:
            UUID(str(actual.pop("id")))
            UUID(str(expected.pop("id")))
        except (KeyError, ValueError):
            return False
        if (
            actual.get("origin") == "user_edited"
            or str(actual.get("contentType", "")).startswith("user_edit")
            or actual.get("sourceLabels") == ["Edited by you"]
        ):
            return False
        expected_review = expected.pop("reviewState", None)
        actual_review = actual.pop("reviewState", None)
        if expected_review == "pending":
            allowed_review_states = {"kept"} if approved else {"pending", "kept"}
            if actual_review not in allowed_review_states:
                return False
        elif actual_review != expected_review:
            return False
        if actual != expected:
            return False
    return True


def _slides_match_deterministic_output(
    actual_slides: list[object],
    expected_slides: list[dict[str, object]],
    *,
    approved: bool,
) -> bool:
    if len(actual_slides) != len(expected_slides):
        return False
    for raw_actual, raw_expected in zip(actual_slides, expected_slides, strict=True):
        if not isinstance(raw_actual, dict):
            return False
        actual = dict(raw_actual)
        expected = dict(raw_expected)
        expected_review = expected.pop("reviewState", None)
        actual_review = actual.pop("reviewState", None)
        if approved and expected_review == "needs_review":
            if actual_review != "ready":
                return False
        elif actual_review != expected_review:
            return False
        if actual != expected:
            return False
    return True


async def _validated_owner_presentation(
    session: AsyncSession,
    tenant: TenantContext,
    service: CreateService,
    company_id: UUID,
    opportunity_id: UUID,
    contact_id: UUID,
    business_case: CreateBusinessCase,
    business_case_version: CreateBusinessCaseVersion,
    *,
    expected_template_id: UUID | None,
    expected_template_version_id: UUID | None,
    storage: VisualStorage | None = None,
) -> CreatePresentation | None:
    presentations = list(
        (
            await session.scalars(
                select(CreatePresentation).where(
                    CreatePresentation.organisation_id == tenant.organisation_id,
                    CreatePresentation.idempotency_key == OWNER_PRESENTATION_KEY,
                )
            )
        ).all()
    )
    if not presentations:
        return None
    if len(presentations) != 1:
        raise _owner_collision("Presentation")
    presentation = presentations[0]
    contact = await session.get(Contact, contact_id)
    company = await session.get(Company, company_id)
    opportunity = await service.repository.opportunity(tenant.organisation_id, opportunity_id)
    template_version = await service.repository.template_version(
        tenant.organisation_id, presentation.template_version_id
    )
    if contact is None or company is None or opportunity is None or template_version is None:
        raise _owner_collision("Presentation")
    expected_audience = [
        {
            "contactId": str(contact_id),
            "name": f"{contact.first_name} {contact.last_name}",
            "role": contact.job_title,
            "audienceType": "finance",
        }
    ]
    context = await service._build_context(
        company,
        opportunity,
        (business_case, business_case_version),
        "all",
    )
    expected_plan = await service._build_plan(
        template_version,
        "business_case",
        context,
        OWNER_PRESENTATION_FOCUS,
    )
    if (
        presentation.account_id != company_id
        or presentation.opportunity_id != opportunity_id
        or presentation.business_case_id != business_case.id
        or presentation.business_case_version_id != business_case_version.id
        or presentation.business_case_scenario != "all"
        or presentation.created_by_user_id != tenant.user_id
        or presentation.title != OWNER_PRESENTATION_TITLE
        or presentation.objective != "business_case"
        or presentation.audience_json != expected_audience
        or presentation.focus_instruction != OWNER_PRESENTATION_FOCUS
        or presentation.source_context_fingerprint != _json_fingerprint(context)
        or _normalise_plan(presentation.plan_json) != _normalise_plan(expected_plan)
        or (expected_template_id is not None and presentation.template_id != expected_template_id)
        or (
            expected_template_version_id is not None
            and presentation.template_version_id != expected_template_version_id
        )
    ):
        raise _owner_collision("Presentation")
    versions = list(
        (
            await session.scalars(
                select(CreatePresentationVersion).where(
                    CreatePresentationVersion.organisation_id == tenant.organisation_id,
                    CreatePresentationVersion.presentation_id == presentation.id,
                )
            )
        ).all()
    )
    if presentation.state == "draft_plan":
        if versions or presentation.review_state != "pending":
            raise _owner_collision("Presentation")
        return presentation
    if len(versions) != 1 or presentation.state not in {"generating", "needs_review", "ready"}:
        raise _owner_collision("Presentation")
    version = versions[0]
    valid_state = (
        (presentation.state == "generating" and version.state == "generating" and version.review_state == "pending")
        or (
            presentation.state == "needs_review"
            and version.state == "needs_review"
            and version.review_state == "pending"
        )
        or (
            presentation.state == "ready"
            and presentation.review_state == "approved"
            and version.state == "ready"
            and version.review_state == "approved"
            and version.approved_by_user_id == tenant.user_id
        )
    )
    if (
        version.version != 1
        or version.template_id != presentation.template_id
        or version.template_version_id != presentation.template_version_id
        or version.created_by_user_id != tenant.user_id
        or version.plan_snapshot_json != presentation.plan_json
        or version.audience_snapshot_json != expected_audience
        or version.source_context_json != context
        or version.source_context_fingerprint != presentation.source_context_fingerprint
        or version.renderer_version != "deterministic_pptx_v1"
        or version.generation_schema_version != 1
        or version.idempotency_key != OWNER_PRESENTATION_GENERATION_KEY
        or not valid_state
    ):
        raise _owner_collision("Presentation")
    if presentation.state == "generating":
        if version.generated_content_json or version.claim_manifest_json:
            raise _owner_collision("Presentation")
        return presentation

    slide_rows = list(
        (
            await session.scalars(
                select(CreateTemplateSlide).where(
                    CreateTemplateSlide.organisation_id == tenant.organisation_id,
                    CreateTemplateSlide.template_version_id == template_version.id,
                )
            )
        ).all()
    )
    content_rows = list(
        (
            await session.scalars(
                select(CreateApprovedContentItem).where(
                    CreateApprovedContentItem.organisation_id == tenant.organisation_id,
                    CreateApprovedContentItem.template_version_id == template_version.id,
                    CreateApprovedContentItem.status == "approved",
                )
            )
        ).all()
    )
    slide_map = {slide.id: slide for slide in slide_rows}
    content_map = {item.slide_id: item for item in content_rows}
    expected_generated, expected_claims = _compose(presentation, version, slide_map, content_map)
    _validate_claim_manifest(expected_generated, expected_claims)
    if not _slides_match_deterministic_output(
        version.generated_content_json,
        expected_generated,
        approved=presentation.state == "ready",
    ) or not _claims_match_deterministic_output(
        version.claim_manifest_json,
        expected_claims,
        approved=presentation.state == "ready",
    ):
        raise _owner_collision("Presentation")
    expected_storage_key = f"{tenant.organisation_id}/create/presentations/{presentation.id}/{version.id}.pptx"
    if (
        version.pptx_storage_key != expected_storage_key
        or version.storage_status != "available"
        or version.byte_size is None
        or version.checksum_sha256 is None
        or version.validation_profile_version is None
        or version.validated_at is None
    ):
        raise _owner_collision("Presentation")
    active_storage = storage or service.storage
    try:
        actual_output = await active_storage.read(expected_storage_key)
        template_source = await active_storage.read(template_version.storage_key)
    except (OSError, VisualStorageError) as exc:
        raise _owner_collision("Presentation") from exc
    if len(actual_output) != version.byte_size or hashlib.sha256(actual_output).hexdigest() != version.checksum_sha256:
        raise _owner_collision("Presentation")
    organisation = await session.get(Organisation, tenant.organisation_id)
    if organisation is None:
        raise _owner_collision("Presentation")
    render_slides = _render_manifest(expected_generated, slide_map, presentation, company, opportunity)
    processor = create_processor(service.settings)
    try:
        expected_output = processor.render(
            template_source,
            tuple(render_slides),
            title=presentation.title,
            organisation_name=organisation.name,
        )
        processor.validate_output(
            actual_output,
            _output_expectations(expected_generated, slide_map, render_slides),
        )
    except (OSError, PptxProcessingError, ValueError) as exc:
        raise _owner_collision("Presentation") from exc
    try:
        actual_canonical = _canonicalise_generated_pptx_archive(actual_output)
        expected_canonical = _canonicalise_generated_pptx_archive(expected_output)
    except (OSError, zipfile.BadZipFile) as exc:
        raise _owner_collision("Presentation") from exc
    if actual_canonical != expected_canonical:
        raise _owner_collision("Presentation")
    return presentation


def _validated_cleanup_storage_keys(organisation_id: UUID, raw_keys: object) -> tuple[str, ...]:
    if not isinstance(raw_keys, list) or len(raw_keys) > OWNER_RESET_CLEANUP_MAX_KEYS:
        raise ValueError("Owner demo reset refused an invalid local object cleanup marker.")
    keys: list[str] = []
    source_ids = demo_source_evidence_ids(organisation_id)
    allowed_base_keys = {
        f"{organisation_id}/documents/{source_ids[f'{label}:source'].hex}.txt" for label in DEMO_DOCUMENTS
    }
    allowed_base_keys.add(f"{organisation_id}/{demo_debrief_ids(organisation_id)[1]}/demo-whiteboard.png")
    for raw_key in raw_keys:
        if not isinstance(raw_key, str) or raw_key in keys:
            raise ValueError("Owner demo reset refused an invalid local object cleanup marker.")
        if raw_key in allowed_base_keys:
            keys.append(raw_key)
            continue
        parts = raw_key.split("/")
        if (
            len(parts) != 5
            or parts[:2] != [str(organisation_id), "create"]
            or parts[2]
            not in {
                "templates",
                "presentations",
            }
        ):
            raise ValueError("Owner demo reset refused an invalid local object cleanup marker.")
        try:
            UUID(parts[3])
            UUID(parts[4].removesuffix(".pptx"))
        except ValueError as exc:
            raise ValueError("Owner demo reset refused an invalid local object cleanup marker.") from exc
        if not parts[4].endswith(".pptx"):
            raise ValueError("Owner demo reset refused an invalid local object cleanup marker.")
        keys.append(raw_key)
    return tuple(keys)


def _cleanup_keys_from_event(event: BetaSystemEvent, organisation_id: UUID) -> tuple[str, ...]:
    expected_id = owner_demo_cleanup_event_id(organisation_id)
    if (
        event.id != expected_id
        or event.organisation_id != organisation_id
        or event.event_type != OWNER_RESET_CLEANUP_EVENT
        or event.subject_id != expected_id
        or set(event.metadata_json) != {"schemaVersion", "storageKeys"}
        or event.metadata_json.get("schemaVersion") != OWNER_RESET_CLEANUP_SCHEMA_VERSION
    ):
        raise ValueError("Owner demo reset refused an invalid local object cleanup marker.")
    return _validated_cleanup_storage_keys(organisation_id, event.metadata_json.get("storageKeys"))


def _canonicalise_pptx_archive(content: bytes) -> bytes:
    """Strip volatile ZIP metadata while preserving the PPTX package content."""

    canonical = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content), "r") as source:
        with zipfile.ZipFile(canonical, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
            for source_info in source.infolist():
                target_info = zipfile.ZipInfo(source_info.filename, date_time=(1980, 1, 1, 0, 0, 0))
                target_info.compress_type = zipfile.ZIP_DEFLATED
                target_info.create_system = source_info.create_system
                target_info.external_attr = source_info.external_attr
                target_info.internal_attr = source_info.internal_attr
                target.writestr(target_info, source.read(source_info.filename))
    return canonical.getvalue()


def _canonicalise_generated_pptx_archive(content: bytes) -> bytes:
    """Normalise container and renderer timestamps for exact output comparison."""

    canonical = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content), "r") as source:
        with zipfile.ZipFile(canonical, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
            for source_info in source.infolist():
                member = source.read(source_info.filename)
                if source_info.filename == "docProps/core.xml":
                    member = re.sub(
                        rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                        rb"\g<1>1980-01-01T00:00:00Z\g<2>",
                        member,
                    )
                target_info = zipfile.ZipInfo(source_info.filename, date_time=(1980, 1, 1, 0, 0, 0))
                target_info.compress_type = zipfile.ZIP_DEFLATED
                target_info.create_system = source_info.create_system
                target_info.external_attr = source_info.external_attr
                target_info.internal_attr = source_info.internal_attr
                target.writestr(target_info, member)
    return canonical.getvalue()


def _template_bytes() -> bytes:
    presentation = Presentation()
    title = presentation.slides.add_slide(presentation.slide_layouts[0])
    title.shapes.title.text = "Northstar multi-site workflow pilot"
    title.placeholders[1].text = "Synthetic owner demonstration · reviewed customer story"
    for heading, body in (
        ("Agenda", "Customer context\nReviewed value case\nProposed approach\nCustomer-owned next steps"),
        (
            "Customer context",
            "Distributed relationship context\nInconsistent follow-through\nA need for a reviewed sales-to-delivery handover",
        ),
        (
            "Proposed approach",
            "Preserve source-labelled context\nReview next actions before use\nCoordinate a customer-owned pilot",
        ),
        ("Business case", "Use only the reviewed synthetic assumptions and outputs shown in the Business Case."),
        ("Next steps", "Confirm security review\nAgree the pilot workshop\nName the implementation lead"),
        (
            "Customer-safe notice",
            "Illustrative synthetic content only. Final scope and commercials require explicit agreement.",
        ),
    ):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = heading
        slide.placeholders[1].text = body
    output = io.BytesIO()
    presentation.save(output)
    # python-pptx writes current ZIP member timestamps. Normalise the container
    # metadata so the same synthetic template has one stable checksum across
    # processes and days, not merely within one seed invocation.
    return _canonicalise_pptx_archive(output.getvalue())


async def _seed_selling_profile(
    session_factory: async_sessionmaker[AsyncSession],
    tenant: TenantContext,
) -> tuple[UUID, UUID, UUID]:
    async with session_factory() as session:
        await set_tenant_database_context(session, tenant.organisation_id)
        owner_revision = await _validated_owner_profile(session, tenant)
        service = SellingProfileService(session, tenant)
        if owner_revision is None:
            management = await service.create_draft(
                SellingProfileDraftCreate(
                    idempotency_key=OWNER_PROFILE_KEY,
                    content=_selling_profile_content(),
                )
            )
            if management.draft is None:
                raise ValueError("Owner demo Company & Selling Profile draft was not created.")
            owner_revision = await _validated_owner_profile(session, tenant)
        if owner_revision is None:
            raise ValueError("Owner demo Company & Selling Profile could not be loaded.")
        if owner_revision.state == "draft":
            management = await service.approve(
                owner_revision.id,
                SellingProfileApproveRequest(expected_lock_version=owner_revision.lock_version),
            )
            if management.current is None:
                raise ValueError("Owner demo Company & Selling Profile was not approved.")
            owner_revision = await _validated_owner_profile(session, tenant)
            if owner_revision is None:
                raise ValueError("Owner demo Company & Selling Profile could not be reloaded.")
            website_source_id, text_source_id = await _seed_company_knowledge_sources(service)
            return owner_revision.id, website_source_id, text_source_id
        if owner_revision.state != "approved":
            raise ValueError("Owner demo Company & Selling Profile is not in an approved state.")
        website_source_id, text_source_id = await _seed_company_knowledge_sources(service)
        return owner_revision.id, website_source_id, text_source_id


async def _seed_business_case(
    session_factory: async_sessionmaker[AsyncSession],
    tenant: TenantContext,
    settings: Settings,
    company_id: UUID,
    opportunity_id: UUID,
) -> tuple[UUID, UUID, UUID]:
    async with session_factory() as session:
        await set_tenant_database_context(session, tenant.organisation_id)
        service = BusinessCaseService(session, tenant, settings)
        validated_model = await _validated_owner_value_model(session, tenant, service)
        model = await service.create_value_model(
            ValueModelCreateRequest(
                name=OWNER_VALUE_MODEL_NAME,
                description=OWNER_VALUE_MODEL_DESCRIPTION,
                definition=_value_model_definition(),
                idempotency_key=OWNER_VALUE_MODEL_KEY,
            )
        )
        validated_model = await _validated_owner_value_model(session, tenant, service)
        if validated_model is None or validated_model[0].id != model.id:
            raise _owner_collision("Value Model")
        if model.latest_version.state == "draft":
            model = await service.approve_value_model(
                model.id,
                model.latest_version.id,
                ValueModelApprovalRequest(confirmed=True),
            )
            validated_model = await _validated_owner_value_model(session, tenant, service)
            if validated_model is None:
                raise _owner_collision("Value Model")
        if model.latest_version.state != "approved":
            raise ValueError("Owner demo Value Model was not approved.")
        model_row, model_version = validated_model
        validated_case = await _validated_owner_business_case(
            session,
            tenant,
            service,
            company_id,
            opportunity_id,
            model_row,
            model_version,
        )
        business_case = await service.create_business_case(
            BusinessCaseCreateRequest(
                account_id=company_id,
                opportunity_id=opportunity_id,
                model_version_id=model.latest_version.id,
                currency="AUD",
                title=OWNER_BUSINESS_CASE_TITLE,
                idempotency_key=OWNER_BUSINESS_CASE_KEY,
            )
        )
        validated_case = await _validated_owner_business_case(
            session,
            tenant,
            service,
            company_id,
            opportunity_id,
            model_row,
            model_version,
        )
        if validated_case is None or validated_case[0].id != business_case.id:
            raise _owner_collision("Business Case")
        if business_case.current_version is None:
            business_case = await service.calculate_business_case(
                business_case.id,
                _business_case_calculation_request(),
            )
            validated_case = await _validated_owner_business_case(
                session,
                tenant,
                service,
                company_id,
                opportunity_id,
                model_row,
                model_version,
            )
            if validated_case is None:
                raise _owner_collision("Business Case")
        if business_case.state == "calculated":
            business_case = await service.approve_business_case(
                business_case.id,
                BusinessCaseApprovalRequest(confirmed=True),
            )
            validated_case = await _validated_owner_business_case(
                session,
                tenant,
                service,
                company_id,
                opportunity_id,
                model_row,
                model_version,
            )
            if validated_case is None:
                raise _owner_collision("Business Case")
        if business_case.state != "approved" or business_case.current_version is None:
            raise ValueError("Owner demo Business Case was not approved.")
        return model.id, business_case.id, business_case.current_version.id


def _create_service(
    session: AsyncSession,
    tenant: TenantContext,
    settings: Settings,
) -> CreateService:
    return CreateService(
        session,
        tenant,
        settings,
        create_visual_storage(settings),
        create_processor(settings),
    )


async def _process_create_work(
    session_factory: async_sessionmaker[AsyncSession],
    settings: Settings,
    organisation_id: UUID,
    work_type: Literal["template", "presentation"],
    work_id: UUID,
    worker_id: str,
) -> None:
    """Claim and execute only the exact owner-demo work item.

    The general worker loop deliberately discovers work across organisations and
    claims the oldest eligible row. A maintenance seed must never use that loop:
    an unrelated queued upload could otherwise be processed as a side effect.
    """

    worker = CreateWorkerService(
        session_factory,
        settings,
        storage=create_visual_storage(settings),
        processor=create_processor(settings),
    )
    now = datetime.now(UTC)
    async with session_factory() as session, session.begin():
        await set_tenant_database_context(session, organisation_id)
        if work_type == "template":
            item = await session.scalar(
                select(CreateTemplateVersion)
                .where(
                    CreateTemplateVersion.organisation_id == organisation_id,
                    CreateTemplateVersion.id == work_id,
                    CreateTemplateVersion.processing_state == "processing",
                    CreateTemplateVersion.processing_attempts < settings.private_beta_create_processing_retries,
                    or_(
                        CreateTemplateVersion.lease_expires_at.is_(None),
                        CreateTemplateVersion.lease_expires_at <= now,
                    ),
                )
                .with_for_update(skip_locked=True)
            )
        else:
            item = await session.scalar(
                select(CreatePresentationVersion)
                .where(
                    CreatePresentationVersion.organisation_id == organisation_id,
                    CreatePresentationVersion.id == work_id,
                    CreatePresentationVersion.state == "generating",
                    CreatePresentationVersion.processing_attempts < settings.private_beta_create_processing_retries,
                    or_(
                        CreatePresentationVersion.lease_expires_at.is_(None),
                        CreatePresentationVersion.lease_expires_at <= now,
                    ),
                )
                .with_for_update(skip_locked=True)
            )
        if item is None:
            raise ValueError(f"Owner demo {work_type} work item is not eligible for an exact claim.")
        item.processing_attempts += 1
        item.worker_id = worker_id
        item.lease_expires_at = now + timedelta(seconds=settings.worker_lease_duration_seconds)
        item.safe_failure_code = None
    claim = ClaimedCreateWork(organisation_id, work_type, work_id, worker_id)
    await worker.execute_claimed(claim)


async def _owner_template(
    session_factory: async_sessionmaker[AsyncSession],
    tenant: TenantContext,
    settings: Settings,
    template_bytes: bytes,
) -> tuple[UUID, TemplateSummaryResponse]:
    """Create or validate the exact deterministic owner-demo template marker."""

    template_id = owner_demo_template_id(tenant.organisation_id)
    checksum = hashlib.sha256(template_bytes).hexdigest()
    async with session_factory() as session:
        await set_tenant_database_context(session, tenant.organisation_id)
        by_id = await session.scalar(
            select(CreateTemplate).where(
                CreateTemplate.organisation_id == tenant.organisation_id,
                CreateTemplate.id == template_id,
            )
        )
        by_name = await session.scalar(
            select(CreateTemplate).where(
                CreateTemplate.organisation_id == tenant.organisation_id,
                CreateTemplate.name == OWNER_TEMPLATE_NAME,
            )
        )
        checksum_version = await session.scalar(
            select(CreateTemplateVersion).where(
                CreateTemplateVersion.organisation_id == tenant.organisation_id,
                CreateTemplateVersion.checksum_sha256 == checksum,
            )
        )
        if by_name is not None and by_name.id != template_id:
            raise ValueError("Owner demo refused: a non-demo template already uses the owner-demo template name.")
        if checksum_version is not None and checksum_version.template_id != template_id:
            raise ValueError("Owner demo refused: a non-demo template already uses the owner-demo template checksum.")
        if by_id is not None and (
            by_id.name != OWNER_TEMPLATE_NAME or by_id.created_by_user_id != tenant.user_id or by_id.state != "active"
        ):
            raise ValueError("Owner demo refused: the deterministic template identity is already owned.")
        if by_id is None:
            by_id = CreateTemplate(
                id=template_id,
                organisation_id=tenant.organisation_id,
                name=OWNER_TEMPLATE_NAME,
                state="active",
                created_by_user_id=tenant.user_id,
            )
            session.add(by_id)
            await session.flush()
        versions = list(
            (
                await session.scalars(
                    select(CreateTemplateVersion)
                    .where(
                        CreateTemplateVersion.organisation_id == tenant.organisation_id,
                        CreateTemplateVersion.template_id == template_id,
                    )
                    .order_by(CreateTemplateVersion.version)
                )
            ).all()
        )
        storage = create_visual_storage(settings)
        for version in versions:
            if (
                version.checksum_sha256 != checksum
                or version.display_filename != OWNER_TEMPLATE_FILE_NAME
                or version.mime_type != PPTX_MIME_TYPE
                or version.byte_size != len(template_bytes)
                or version.uploaded_by_user_id != tenant.user_id
                or version.authority_attested_by_user_id != tenant.user_id
            ):
                raise ValueError("Owner demo refused: the deterministic template contains non-demo content.")
            try:
                stored = await storage.read(version.storage_key)
                canonical_checksum = hashlib.sha256(_canonicalise_pptx_archive(stored)).hexdigest()
            except (OSError, VisualStorageError, zipfile.BadZipFile) as exc:
                raise ValueError("Owner demo refused: the deterministic template blob could not be verified.") from exc
            if (
                len(stored) != version.byte_size
                or hashlib.sha256(stored).hexdigest() != version.checksum_sha256
                or canonical_checksum != checksum
            ):
                raise ValueError("Owner demo refused: the deterministic template contains non-demo content.")
        service = _create_service(session, tenant, settings)
        if not versions:
            template = await service.upload_template(
                TemplateUploadRequest(
                    template_id=template_id,
                    name=OWNER_TEMPLATE_NAME,
                    file_name=OWNER_TEMPLATE_FILE_NAME,
                    mime_type=PPTX_MIME_TYPE,
                    content_base64=base64.b64encode(template_bytes).decode("ascii"),
                    checksum_sha256=checksum,
                    authority_attested=True,
                    attestation_version=1,
                )
            )
        else:
            template = await service.get_template(template_id)
        return template_id, template


async def _seed_presentation(
    session_factory: async_sessionmaker[AsyncSession],
    tenant: TenantContext,
    settings: Settings,
    company_id: UUID,
    opportunity_id: UUID,
    contact_id: UUID,
    business_case_version_id: UUID,
) -> tuple[UUID, UUID, UUID]:
    async with session_factory() as session:
        await set_tenant_database_context(session, tenant.organisation_id)
        business_case_version = await session.get(CreateBusinessCaseVersion, business_case_version_id)
        business_case = (
            await session.get(CreateBusinessCase, business_case_version.case_id)
            if business_case_version is not None
            else None
        )
        if business_case is None or business_case_version is None:
            raise ValueError("Owner demo Business Case could not be loaded for the presentation.")
        service = _create_service(session, tenant, settings)
        await _validated_owner_presentation(
            session,
            tenant,
            service,
            company_id,
            opportunity_id,
            contact_id,
            business_case,
            business_case_version,
            expected_template_id=owner_demo_template_id(tenant.organisation_id),
            expected_template_version_id=None,
        )
    template_bytes = _template_bytes()
    template_id, template = await _owner_template(session_factory, tenant, settings, template_bytes)
    if template.latest_version.processing_state == "processing":
        await _process_create_work(
            session_factory,
            settings,
            tenant.organisation_id,
            "template",
            template.latest_version.id,
            "owner-demo-template",
        )
    async with session_factory() as session:
        await set_tenant_database_context(session, tenant.organisation_id)
        service = _create_service(session, tenant, settings)
        template = await service.get_template(template_id)
        if template.latest_version.processing_state not in {"ready", "partial"}:
            raise ValueError("Owner demo presentation template did not process successfully.")
        if template.latest_version.approval_state != "approved":
            for slide in template.latest_version.slides:
                if slide.reuse_state != "pending":
                    continue
                locked = slide.category == "appendix"
                template = await service.update_slide(
                    slide.id,
                    TemplateSlideUpdate(
                        category=slide.category,
                        reuse_state="approved",
                        modification_policy="locked" if locked else "text_placeholders_only",
                        customer_safe=True,
                        required=locked or slide.category == "title",
                        exact_text_required=locked,
                        approved_description="Approved synthetic content for the private owner demonstration.",
                        placeholder_mappings={},
                    ),
                )
            template = await service.approve_template(
                template.id,
                template.latest_version.id,
                TemplateApprovalRequest(confirmed=True),
            )
        presentation = await service.create_presentation(
            PresentationBriefRequest(
                account_id=company_id,
                opportunity_id=opportunity_id,
                objective="business_case",
                audience=[
                    PresentationAudienceInput(
                        contact_id=contact_id,
                        audience_type="finance",
                    )
                ],
                template_version_id=template.latest_version.id,
                business_case_version_id=business_case_version_id,
                business_case_scenario="all",
                focus_instruction=OWNER_PRESENTATION_FOCUS,
                title=OWNER_PRESENTATION_TITLE,
                idempotency_key=OWNER_PRESENTATION_KEY,
            )
        )
        business_case_version = await session.get(CreateBusinessCaseVersion, business_case_version_id)
        business_case = (
            await session.get(CreateBusinessCase, business_case_version.case_id)
            if business_case_version is not None
            else None
        )
        if business_case is None or business_case_version is None:
            raise ValueError("Owner demo Business Case could not be reloaded for the presentation.")
        validated_presentation = await _validated_owner_presentation(
            session,
            tenant,
            service,
            company_id,
            opportunity_id,
            contact_id,
            business_case,
            business_case_version,
            expected_template_id=template_id,
            expected_template_version_id=template.latest_version.id,
        )
        if validated_presentation is None or validated_presentation.id != presentation.id:
            raise _owner_collision("Presentation")
        presentation_id = presentation.id
        template_version_id = template.latest_version.id
        if presentation.current_version is None:
            presentation = await service.generate(
                presentation.id,
                PresentationGenerateRequest(idempotency_key=OWNER_PRESENTATION_GENERATION_KEY),
            )
    if presentation.state == "generating":
        if presentation.current_version is None:
            raise ValueError("Owner demo presentation did not expose its queued version.")
        await _process_create_work(
            session_factory,
            settings,
            tenant.organisation_id,
            "presentation",
            presentation.current_version.id,
            "owner-demo-presentation",
        )
    async with session_factory() as session:
        await set_tenant_database_context(session, tenant.organisation_id)
        service = _create_service(session, tenant, settings)
        presentation = await service.get_presentation(presentation_id)
        business_case_version = await session.get(CreateBusinessCaseVersion, business_case_version_id)
        business_case = (
            await session.get(CreateBusinessCase, business_case_version.case_id)
            if business_case_version is not None
            else None
        )
        if business_case is None or business_case_version is None:
            raise ValueError("Owner demo Business Case could not be reloaded for presentation review.")
        validated_presentation = await _validated_owner_presentation(
            session,
            tenant,
            service,
            company_id,
            opportunity_id,
            contact_id,
            business_case,
            business_case_version,
            expected_template_id=template_id,
            expected_template_version_id=template_version_id,
        )
        if validated_presentation is None:
            raise _owner_collision("Presentation")
        if presentation.state == "needs_review" and presentation.current_version is not None:
            pending = [
                PresentationClaimDecision(claim_id=claim.id, action="keep")
                for claim in presentation.current_version.claims
                if claim.review_state == "pending"
            ]
            if pending:
                presentation = await service.review_claims(
                    presentation.id,
                    PresentationReviewRequest(decisions=pending),
                )
            presentation = await service.approve_presentation(
                presentation.id,
                PresentationApprovalRequest(confirmed=True),
            )
            validated_presentation = await _validated_owner_presentation(
                session,
                tenant,
                service,
                company_id,
                opportunity_id,
                contact_id,
                business_case,
                business_case_version,
                expected_template_id=template_id,
                expected_template_version_id=template_version_id,
            )
            if validated_presentation is None:
                raise _owner_collision("Presentation")
        if presentation.state != "ready" or presentation.current_version is None:
            raise ValueError("Owner demo presentation was not approved and rendered.")
        return template_version_id, presentation.id, presentation.current_version.id


def _owner_deal_room_content(
    organisation_id: UUID,
    opportunity_id: UUID,
    contact_id: UUID,
    business_case_version_id: UUID,
    presentation_version_id: UUID,
) -> DealRoomDraftContent:
    campaign_ids = demo_campaign_ids(organisation_id)
    return DealRoomDraftContent(
        overview=(
            "Synthetic collaboration space for Northstar's reviewed multi-site workflow pilot. "
            "No real customer data or live provider output is included."
        ),
        commercial_summary=(
            "Illustrative AUD Business Case only. Final scope and commercials require explicit agreement."
        ),
        business_case_version_id=business_case_version_id,
        stakeholders=[
            DealRoomStakeholderDraft(
                id=uuid.uuid5(opportunity_id, "owner-demo-stakeholder-jane"),
                name="[DEMO] Jane Smith",
                role="Chief Information Officer",
                company="[DEMO] Northstar Facilities Group",
                party="customer",
                source_contact_id=campaign_ids["contact-jane"],
            ),
            DealRoomStakeholderDraft(
                id=uuid.uuid5(opportunity_id, "owner-demo-stakeholder-jordan"),
                name="[DEMO] Jordan Lee",
                role="Finance approver",
                company="[DEMO] Northstar Facilities Group",
                party="customer",
                source_contact_id=contact_id,
            ),
        ],
        milestones=[
            DealRoomMilestoneDraft(
                id=uuid.uuid5(opportunity_id, "owner-demo-milestone-security"),
                title="Complete the synthetic security review",
                owner_party=DealRoomOwnerParty.OUR_TEAM,
                target_date=date(2026, 11, 3),
                status=DealRoomMilestoneStatus.IN_PROGRESS,
                note="Review the prepared security summary before any customer use.",
            ),
            DealRoomMilestoneDraft(
                id=uuid.uuid5(opportunity_id, "owner-demo-milestone-pilot"),
                title="Agree the pilot workshop",
                owner_party=DealRoomOwnerParty.CUSTOMER,
                target_date=date(2026, 11, 10),
                status=DealRoomMilestoneStatus.NOT_STARTED,
            ),
        ],
        resources=[
            DealRoomResourceDraft(
                id=uuid.uuid5(opportunity_id, "owner-demo-resource-presentation"),
                kind=DealRoomResourceKind.PRESENTATION,
                title="Northstar workflow pilot presentation",
                presentation_version_id=presentation_version_id,
            )
        ],
        next_meeting_at=datetime(2026, 11, 5, 2, 0, tzinfo=UTC),
    )


async def _seed_owner_task(
    session_factory: async_sessionmaker[AsyncSession],
    tenant: TenantContext,
    company_id: UUID,
    opportunity_id: UUID,
    contact_id: UUID,
) -> UUID:
    task_id = owner_demo_task_id(tenant.organisation_id)
    expected_due_at = datetime(2026, 11, 4, 2, 0, tzinfo=UTC)
    async with session_factory() as session, session.begin():
        await set_tenant_database_context(session, tenant.organisation_id)
        task = await session.scalar(
            select(Task)
            .where(
                Task.organisation_id == tenant.organisation_id,
                Task.id == task_id,
            )
            .with_for_update()
        )
        if task is None:
            session.add(
                Task(
                    id=task_id,
                    organisation_id=tenant.organisation_id,
                    company_id=company_id,
                    contact_id=contact_id,
                    opportunity_id=opportunity_id,
                    title="[DEMO] Prepare the Northstar security workshop",
                    description=(
                        "Synthetic owner-demo follow-up. Confirm the reviewed security summary and workshop agenda."
                    ),
                    status="open",
                    priority="high",
                    due_at=expected_due_at,
                    assigned_user_id=tenant.user_id,
                    created_by_user_id=tenant.user_id,
                )
            )
            return task_id
        due_at = task.due_at
        if due_at is not None and due_at.utcoffset() is None:
            due_at = due_at.replace(tzinfo=UTC)
        if (
            task.company_id != company_id
            or task.contact_id != contact_id
            or task.opportunity_id != opportunity_id
            or task.title != "[DEMO] Prepare the Northstar security workshop"
            or task.description
            != "Synthetic owner-demo follow-up. Confirm the reviewed security summary and workshop agenda."
            or task.status != "open"
            or task.priority != "high"
            or due_at != expected_due_at
            or task.assigned_user_id != tenant.user_id
            or task.created_by_user_id != tenant.user_id
        ):
            raise _owner_collision("Task")
        return task.id


async def _validated_owner_deal_room(
    session: AsyncSession,
    tenant: TenantContext,
    settings: Settings,
    opportunity_id: UUID,
    expected_content: DealRoomDraftContent,
    *,
    allow_owner_lifecycle: bool = False,
) -> tuple[DealRoom, Literal["create_only", "updated_draft", "published"]] | None:
    room = await session.scalar(
        select(DealRoom).where(
            DealRoom.organisation_id == tenant.organisation_id,
            DealRoom.opportunity_id == opportunity_id,
        )
    )
    if room is None:
        return None
    revisions = list(
        (
            await session.scalars(
                select(DealRoomRevision)
                .where(
                    DealRoomRevision.organisation_id == tenant.organisation_id,
                    DealRoomRevision.room_id == room.id,
                )
                .order_by(DealRoomRevision.revision)
            )
        ).all()
    )
    expected_json = expected_content.model_dump(mode="json", by_alias=True)
    if room.created_by_user_id != tenant.user_id:
        raise _owner_collision("Deal Room")
    links = list(
        (
            await session.scalars(
                select(DealRoomAccessLink).where(
                    DealRoomAccessLink.organisation_id == tenant.organisation_id,
                    DealRoomAccessLink.room_id == room.id,
                )
            )
        ).all()
    )
    audits = list(
        (
            await session.scalars(
                select(DealRoomAuditEvent).where(
                    DealRoomAuditEvent.organisation_id == tenant.organisation_id,
                    DealRoomAuditEvent.room_id == room.id,
                )
            )
        ).all()
    )
    allowed_actions = {
        "created",
        "draft_updated",
        "resource_changed",
        "publication_source_changed",
        "published",
        "republished",
        "paused",
        "revoked",
        "link_rotated",
    }
    if any(audit.actor_user_id != tenant.user_id or audit.action not in allowed_actions for audit in audits):
        raise _owner_collision("Deal Room")
    action_counts = {action: sum(audit.action == action for audit in audits) for action in allowed_actions}
    if action_counts["created"] != 1:
        raise _owner_collision("Deal Room")
    for link in links:
        if (
            link.created_by_user_id != tenant.user_id
            or link.revoked_by_user_id not in {None, tenant.user_id}
            or len(link.token_hash) != 64
            or link.token_hash != link.token_hash.casefold()
        ):
            raise _owner_collision("Deal Room")
    active_links = [link for link in links if link.revoked_at is None]
    if len(active_links) > 1:
        raise _owner_collision("Deal Room")

    empty_json = DealRoomDraftContent().model_dump(mode="json", by_alias=True)
    if not revisions:
        if (
            links
            or room.status != "draft"
            or room.published_revision_id is not None
            or room.last_published_at is not None
        ):
            raise _owner_collision("Deal Room")
        if room.paused_at is not None or room.revoked_at is not None:
            raise _owner_collision("Deal Room")
        if (
            room.draft_version == 1
            and room.lock_version == 1
            and room.draft_content_json == empty_json
            and len(audits) == 1
        ):
            return room, "create_only"
        if (
            room.draft_version == 2
            and room.lock_version == 2
            and room.draft_content_json == expected_json
            and action_counts["draft_updated"] == 1
            and action_counts["resource_changed"] == 1
            and action_counts["publication_source_changed"] == 1
            and len(audits) == 4
        ):
            return room, "updated_draft"
        raise _owner_collision("Deal Room")

    if (
        room.draft_version != 2
        or room.draft_content_json != expected_json
        or room.published_revision_id != revisions[-1].id
        or room.last_published_at is None
        or action_counts["draft_updated"] != 1
        or action_counts["resource_changed"] != 1
        or action_counts["publication_source_changed"] != 1
        or action_counts["published"] != 1
        or action_counts["republished"] != len(revisions) - 1
        or len(links) != action_counts["link_rotated"] + 1
    ):
        raise _owner_collision("Deal Room")
    if not allow_owner_lifecycle and (
        room.status != "published"
        or room.lock_version != 3
        or room.paused_at is not None
        or room.revoked_at is not None
        or len(revisions) != 1
        or action_counts["link_rotated"]
        or len(links) != 1
    ):
        raise _owner_collision("Deal Room")
    if allow_owner_lifecycle:
        mutation_count = sum(
            action_counts[action]
            for action in ("draft_updated", "published", "republished", "paused", "revoked", "link_rotated")
        )
        if room.status not in {"published", "paused", "revoked"} or room.lock_version != 1 + mutation_count:
            raise _owner_collision("Deal Room")
        if room.status == "published" and len(active_links) != 1:
            raise _owner_collision("Deal Room")
        if room.status == "paused" and (room.paused_at is None or room.revoked_at is not None):
            raise _owner_collision("Deal Room")
        if room.status == "revoked" and (room.revoked_at is None or active_links):
            raise _owner_collision("Deal Room")

    service = DealRoomService(session, tenant, settings)
    opportunity = await service.repository.opportunity(tenant.organisation_id, opportunity_id)
    if opportunity is None:
        raise _owner_collision("Deal Room")
    for index, revision in enumerate(revisions, start=1):
        if (
            revision.revision != index
            or revision.snapshot_schema_version != 1
            or revision.published_by_user_id != tenant.user_id
        ):
            raise _owner_collision("Deal Room")
        published_at = revision.published_at
        if published_at.utcoffset() is None:
            published_at = published_at.replace(tzinfo=UTC)
        expected_snapshot = await service._snapshot(opportunity, expected_content, index, published_at)
        canonical = json.dumps(expected_snapshot, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
        if (
            revision.snapshot_json != expected_snapshot
            or revision.content_fingerprint != hashlib.sha256(canonical.encode("utf-8")).hexdigest()
        ):
            raise _owner_collision("Deal Room")
    return room, "published"


async def _seed_deal_room(
    session_factory: async_sessionmaker[AsyncSession],
    tenant: TenantContext,
    settings: Settings,
    opportunity_id: UUID,
    contact_id: UUID,
    business_case_version_id: UUID,
    presentation_version_id: UUID,
) -> UUID:
    content = _owner_deal_room_content(
        tenant.organisation_id,
        opportunity_id,
        contact_id,
        business_case_version_id,
        presentation_version_id,
    )
    async with session_factory() as session:
        await set_tenant_database_context(session, tenant.organisation_id)
        service = DealRoomService(session, tenant, settings)
        validated = await _validated_owner_deal_room(session, tenant, settings, opportunity_id, content)
        if validated is None:
            workspace = (await service.create(opportunity_id)).workspace
            if workspace.room is None:
                raise ValueError("Owner demo Deal Room was not created.")
            room_id = workspace.room.id
            draft_version = workspace.room.draft_version
            lock_version = workspace.room.lock_version
            stage: Literal["create_only", "updated_draft", "published"] = "create_only"
        else:
            persisted_room, stage = validated
            room_id = persisted_room.id
            draft_version = persisted_room.draft_version
            lock_version = persisted_room.lock_version
        if stage == "published":
            return room_id
        if stage == "create_only":
            mutation = await service.update_draft(
                opportunity_id,
                DealRoomDraftUpdate(
                    expected_draft_version=draft_version,
                    expected_lock_version=lock_version,
                    content=content,
                ),
            )
            updated_room = mutation.workspace.room
            if updated_room is None:
                raise ValueError("Owner demo Deal Room draft was not saved.")
            draft_version = updated_room.draft_version
            lock_version = updated_room.lock_version
        mutation = await service.publish(
            opportunity_id,
            DealRoomPublishRequest(
                expected_draft_version=draft_version,
                expected_lock_version=lock_version,
                confirmed=True,
            ),
        )
        published_room = mutation.workspace.room
        if published_room is None:
            raise ValueError("Owner demo Deal Room publication was not returned.")
        validated = await _validated_owner_deal_room(session, tenant, settings, opportunity_id, content)
        if validated is None or validated[1] != "published":
            raise _owner_collision("Deal Room")
        return validated[0].id


def _normalised_handover_content(
    raw_content: dict[str, object],
    source_id_mapping: dict[str, str] | None = None,
) -> dict[str, object]:
    try:
        content = HandoverContent.model_validate(raw_content).model_dump(mode="json", by_alias=True)
    except ValueError as exc:
        raise _owner_collision("Closed-Won Handover") from exc
    seen_item_ids: set[UUID] = set()
    for section in SECTION_KEYS:
        section_alias = section.split("_")[0] + "".join(part.title() for part in section.split("_")[1:])
        raw_items = content.get(section_alias)
        if not isinstance(raw_items, list):
            raise _owner_collision("Closed-Won Handover")
        for raw_item in raw_items:
            if not isinstance(raw_item, dict):
                raise _owner_collision("Closed-Won Handover")
            item_id = raw_item.pop("id", None)
            try:
                parsed_item_id = UUID(str(item_id))
            except ValueError as exc:
                raise _owner_collision("Closed-Won Handover") from exc
            if parsed_item_id in seen_item_ids:
                raise _owner_collision("Closed-Won Handover")
            seen_item_ids.add(parsed_item_id)
            if source_id_mapping is not None:
                source_ids = raw_item.get("sourceIds")
                if not isinstance(source_ids, list):
                    raise _owner_collision("Closed-Won Handover")
                try:
                    raw_item["sourceIds"] = [source_id_mapping[str(source_id)] for source_id in source_ids]
                except KeyError as exc:
                    raise _owner_collision("Closed-Won Handover") from exc
    return content


async def _validated_owner_handover(
    session: AsyncSession,
    tenant: TenantContext,
    service: HandoverService,
    opportunity_id: UUID,
    *,
    allow_owner_lifecycle: bool = False,
) -> tuple[ClosedWonHandover, ClosedWonHandoverRevision] | None:
    handover = await session.scalar(
        select(ClosedWonHandover).where(
            ClosedWonHandover.organisation_id == tenant.organisation_id,
            ClosedWonHandover.opportunity_id == opportunity_id,
        )
    )
    if handover is None:
        return None
    revisions = list(
        (
            await session.scalars(
                select(ClosedWonHandoverRevision)
                .where(
                    ClosedWonHandoverRevision.organisation_id == tenant.organisation_id,
                    ClosedWonHandoverRevision.handover_id == handover.id,
                )
                .order_by(ClosedWonHandoverRevision.revision)
            )
        ).all()
    )
    if (
        allow_owner_lifecycle
        and handover.created_by_user_id == tenant.user_id
        and handover.lock_version == 1
        and len(revisions) == 1
        and revisions[0].status == "draft"
        and revisions[0].lock_version == 1
    ):
        return await _validated_owner_handover(
            session,
            tenant,
            service,
            opportunity_id,
            allow_owner_lifecycle=False,
        )
    if allow_owner_lifecycle:
        if handover.created_by_user_id != tenant.user_id or not revisions:
            raise _owner_collision("Closed-Won Handover")
        revision_ids = {revision.id for revision in revisions}
        audits = list(
            (
                await session.scalars(
                    select(ClosedWonHandoverAuditEvent).where(
                        ClosedWonHandoverAuditEvent.organisation_id == tenant.organisation_id,
                        ClosedWonHandoverAuditEvent.handover_id == handover.id,
                    )
                )
            ).all()
        )
        allowed_actions = {
            "draft_created",
            "draft_updated",
            "submitted_for_review",
            "claim_confirmed",
            "approved",
            "superseded",
            "retired",
        }
        if len(audits) != handover.lock_version or any(
            audit.actor_user_id != tenant.user_id
            or audit.revision_id not in revision_ids
            or audit.action not in allowed_actions
            for audit in audits
        ):
            raise _owner_collision("Closed-Won Handover")
        for index, revision in enumerate(revisions, start=1):
            revision_audits = [audit for audit in audits if audit.revision_id == revision.id]
            if (
                revision.revision != index
                or revision.content_schema_version != 1
                or revision.created_by_user_id != tenant.user_id
                or revision.lock_version != len(revision_audits)
                or sum(audit.action == "draft_created" for audit in revision_audits) != 1
                or revision.status not in {"draft", "in_review", "approved", "superseded", "retired"}
                or revision.submitted_by_user_id not in {None, tenant.user_id}
                or revision.approved_by_user_id not in {None, tenant.user_id}
                or revision.retired_by_user_id not in {None, tenant.user_id}
            ):
                raise _owner_collision("Closed-Won Handover")
            if revision.status == "draft" and any(
                value is not None
                for value in (
                    revision.submitted_by_user_id,
                    revision.submitted_at,
                    revision.approved_by_user_id,
                    revision.approved_at,
                )
            ):
                raise _owner_collision("Closed-Won Handover")
            if revision.status == "in_review" and (
                revision.submitted_by_user_id != tenant.user_id
                or revision.submitted_at is None
                or revision.approved_by_user_id is not None
                or revision.approved_at is not None
            ):
                raise _owner_collision("Closed-Won Handover")
            if revision.status in {"approved", "superseded", "retired"} and (
                revision.submitted_by_user_id != tenant.user_id
                or revision.submitted_at is None
                or revision.approved_by_user_id != tenant.user_id
                or revision.approved_at is None
            ):
                raise _owner_collision("Closed-Won Handover")
            sources = list(
                (
                    await session.scalars(
                        select(ClosedWonHandoverSource).where(
                            ClosedWonHandoverSource.organisation_id == tenant.organisation_id,
                            ClosedWonHandoverSource.revision_id == revision.id,
                        )
                    )
                ).all()
            )
            if not sources or len(revision.source_pack_fingerprint) != 64:
                raise _owner_collision("Closed-Won Handover")
            definitions: list[SourceDefinition] = []
            source_ids: set[UUID] = set()
            for source in sources:
                try:
                    source_type = HandoverSourceType(source.source_type)
                    authority_type = HandoverAuthorityType(source.authority_type)
                except ValueError as exc:
                    raise _owner_collision("Closed-Won Handover") from exc
                if (
                    source.opportunity_id != opportunity_id
                    or not source.label.strip()
                    or len(source.source_fingerprint) != 64
                    or source.source_fingerprint != source.source_fingerprint.casefold()
                ):
                    raise _owner_collision("Closed-Won Handover")
                source_ids.add(source.id)
                definitions.append(
                    SourceDefinition(
                        reference_id=source.id,
                        source_type=source_type,
                        source_id=source.source_id,
                        source_version_id=source.source_version_id,
                        source_version=source.source_version,
                        authority_type=authority_type,
                        label=source.label,
                        snapshot=source.snapshot_json,
                        fingerprint=source.source_fingerprint,
                    )
                )
            try:
                content = HandoverContent.model_validate(revision.content_json)
            except ValueError as exc:
                raise _owner_collision("Closed-Won Handover") from exc
            for section in SECTION_KEYS:
                for item in getattr(content, section):
                    if not set(item.source_ids).issubset(source_ids) or item.confirmed_by_user_id not in {
                        None,
                        tenant.user_id,
                    }:
                        raise _owner_collision("Closed-Won Handover")
            content_was_mutated = any(audit.action in {"draft_updated", "claim_confirmed"} for audit in revision_audits)
            if not content_was_mutated:
                source_opportunity = next(
                    (
                        definition
                        for definition in definitions
                        if definition.source_type == HandoverSourceType.OPPORTUNITY
                    ),
                    None,
                )
                opportunity = await service.repository.opportunity(tenant.organisation_id, opportunity_id)
                if opportunity is None or source_opportunity is None:
                    raise _owner_collision("Closed-Won Handover")
                snapshot = source_opportunity.snapshot
                frozen_opportunity = copy.copy(opportunity)
                frozen_opportunity.name = str(snapshot.get("name"))
                estimated_value = snapshot.get("estimatedValue")
                frozen_opportunity.estimated_value = (
                    Decimal(estimated_value) if isinstance(estimated_value, str) else None
                )
                frozen_opportunity.currency = cast(str | None, snapshot.get("currency"))
                frozen_opportunity.status = cast(str, snapshot.get("status"))
                actual_close_date = snapshot.get("actualCloseDate")
                frozen_opportunity.actual_close_date = (
                    date.fromisoformat(actual_close_date) if isinstance(actual_close_date, str) else None
                )
                expected = service._initial_content(frozen_opportunity, definitions).model_dump(
                    mode="json", by_alias=True
                )
                if _normalised_handover_content(revision.content_json) != _normalised_handover_content(expected):
                    raise _owner_collision("Closed-Won Handover")
        return handover, revisions[-1]
    if handover.created_by_user_id != tenant.user_id or handover.lock_version != 1 or len(revisions) != 1:
        raise _owner_collision("Closed-Won Handover")
    revision = revisions[0]
    if (
        revision.opportunity_id != opportunity_id
        or revision.revision != 1
        or revision.status != "draft"
        or revision.content_schema_version != 1
        or revision.lock_version != 1
        or revision.created_by_user_id != tenant.user_id
        or revision.submitted_by_user_id is not None
        or revision.submitted_at is not None
        or revision.approved_by_user_id is not None
        or revision.approved_at is not None
        or revision.superseded_at is not None
        or revision.retired_by_user_id is not None
        or revision.retired_at is not None
        or revision.retirement_reason is not None
    ):
        raise _owner_collision("Closed-Won Handover")
    opportunity = await service.repository.opportunity(tenant.organisation_id, opportunity_id)
    if opportunity is None:
        raise _owner_collision("Closed-Won Handover")
    definitions = await service._source_definitions(opportunity)
    sources = list(
        (
            await session.scalars(
                select(ClosedWonHandoverSource).where(
                    ClosedWonHandoverSource.organisation_id == tenant.organisation_id,
                    ClosedWonHandoverSource.revision_id == revision.id,
                )
            )
        ).all()
    )
    source_by_key = {
        (source.source_type, source.source_id, source.source_version_id, source.source_version): source
        for source in sources
    }
    if len(source_by_key) != len(sources) or len(sources) != len(definitions):
        raise _owner_collision("Closed-Won Handover")
    source_id_mapping: dict[str, str] = {}
    for definition in definitions:
        key = (
            definition.source_type.value,
            definition.source_id,
            definition.source_version_id,
            definition.source_version,
        )
        matched_source = source_by_key.get(key)
        if matched_source is None or (
            matched_source.opportunity_id != opportunity_id
            or matched_source.authority_type != definition.authority_type.value
            or matched_source.label != definition.label
            or matched_source.snapshot_json != definition.snapshot
            or matched_source.source_fingerprint != definition.fingerprint
        ):
            raise _owner_collision(f"Closed-Won Handover source {key[0]}:{key[1]}")
        source_id_mapping[str(definition.reference_id)] = str(matched_source.id)
    expected_content = service._initial_content(opportunity, definitions).model_dump(mode="json", by_alias=True)
    if revision.source_pack_fingerprint != service._source_pack_fingerprint(
        definitions
    ) or _normalised_handover_content(revision.content_json) != _normalised_handover_content(
        expected_content, source_id_mapping
    ):
        raise _owner_collision("Closed-Won Handover")
    return handover, revision


async def _seed_handover(
    session_factory: async_sessionmaker[AsyncSession],
    tenant: TenantContext,
    settings: Settings,
    opportunity_id: UUID,
) -> tuple[UUID, UUID]:
    async with session_factory() as session:
        await set_tenant_database_context(session, tenant.organisation_id)
        service = HandoverService(session, tenant, settings)
        existing = await _validated_owner_handover(session, tenant, service, opportunity_id)
        if existing is not None:
            return existing[0].id, existing[1].id
        await service.prepare(opportunity_id)
        validated = await _validated_owner_handover(session, tenant, service, opportunity_id)
        if validated is None:
            raise ValueError("Owner demo Closed-Won Handover draft was not prepared.")
        return validated[0].id, validated[1].id


async def seed_owner_demo_data(
    session_factory: async_sessionmaker[AsyncSession],
    organisation_id: UUID,
    user_id: UUID,
    settings: Settings | None = None,
) -> dict[str, object]:
    active_settings = settings or get_settings()
    _assert_private_demo_settings(active_settings)
    _assert_private_demo_database_binding(active_settings, session_factory)
    tenant = await _require_owner_demo_admin(session_factory, organisation_id, user_id)
    await _assert_base_demo_fixture_safe(session_factory, organisation_id, user_id)
    base = await _owner_base_fixture(session_factory, tenant, active_settings)
    company_id, opportunity_id, _, _ = demo_ids(organisation_id)
    contact_id = demo_phone_contact_id(organisation_id)
    (
        selling_profile_revision_id,
        knowledge_website_source_id,
        knowledge_text_source_id,
    ) = await _seed_selling_profile(session_factory, tenant)
    value_model_id, business_case_id, business_case_version_id = await _seed_business_case(
        session_factory,
        tenant,
        active_settings,
        company_id,
        opportunity_id,
    )
    template_version_id, presentation_id, presentation_version_id = await _seed_presentation(
        session_factory,
        tenant,
        active_settings,
        company_id,
        opportunity_id,
        contact_id,
        business_case_version_id,
    )
    deal_room_id = await _seed_deal_room(
        session_factory,
        tenant,
        active_settings,
        opportunity_id,
        contact_id,
        business_case_version_id,
        presentation_version_id,
    )
    task_id = await _seed_owner_task(
        session_factory,
        tenant,
        company_id,
        opportunity_id,
        contact_id,
    )
    handover_id, handover_revision_id = await _seed_handover(
        session_factory,
        tenant,
        active_settings,
        opportunity_id,
    )
    extensions = {
        "selling_profile_revision_id": selling_profile_revision_id,
        "knowledge_website_source_id": knowledge_website_source_id,
        "knowledge_text_source_id": knowledge_text_source_id,
        "value_model_id": value_model_id,
        "business_case_id": business_case_id,
        "business_case_version_id": business_case_version_id,
        "template_version_id": template_version_id,
        "presentation_id": presentation_id,
        "presentation_version_id": presentation_version_id,
        "deal_room_id": deal_room_id,
        "handover_id": handover_id,
        "handover_revision_id": handover_revision_id,
        "task_id": task_id,
    }
    await _finalise_owner_manifest(session_factory, tenant, extensions)
    return {
        **base,
        "dataset_version": OWNER_DEMO_DATASET_VERSION,
        **extensions,
        "provider_calls": 0,
    }


async def reset_owner_demo_extensions(
    session_factory: async_sessionmaker[AsyncSession],
    organisation_id: UUID,
    settings: Settings | None = None,
    *,
    database_finalize: Callable[[AsyncSession], Awaitable[Sequence[str]]] | None = None,
) -> None:
    active_settings = settings or get_settings()
    _assert_private_demo_settings(active_settings)
    _assert_private_demo_database_binding(active_settings, session_factory)
    preflight_tenant = await _owner_demo_reset_tenant(session_factory, organisation_id)
    await _assert_base_demo_fixture_safe(
        session_factory,
        organisation_id,
        allow_owner_lifecycle=True,
    )
    company_id, opportunity_id, _, _ = demo_ids(organisation_id)
    contact_id = demo_phone_contact_id(organisation_id)
    storage = create_visual_storage(active_settings)
    cleanup_event_id = owner_demo_cleanup_event_id(organisation_id)
    cleanup_keys: tuple[str, ...] = ()
    async with session_factory() as session, session.begin():
        await set_tenant_database_context(session, organisation_id)
        if session.get_bind().dialect.name == "postgresql":
            await session.execute(text("SELECT set_config('app.beta_maintenance', 'approved', true)"))
        cleanup_event = await session.scalar(
            select(BetaSystemEvent)
            .where(
                BetaSystemEvent.organisation_id == organisation_id,
                BetaSystemEvent.id == cleanup_event_id,
            )
            .with_for_update()
        )
        pending_cleanup_keys = (
            _cleanup_keys_from_event(cleanup_event, organisation_id) if cleanup_event is not None else ()
        )
        tenant = preflight_tenant
        if tenant is None and cleanup_event is not None:
            if cleanup_event.actor_user_id is None:
                raise ValueError("Owner demo reset refused an unauthorised local object cleanup marker.")
            membership = await session.get(
                OrganisationMembership,
                (organisation_id, cleanup_event.actor_user_id),
            )
            if membership is None or membership.status != "active" or membership.role != "admin":
                raise ValueError("Owner demo reset refused an unauthorised local object cleanup marker.")
            tenant = TenantContext(
                organisation_id=organisation_id,
                user_id=cleanup_event.actor_user_id,
                role=cast(Role, membership.role),
            )
        owner_manifest = await session.scalar(
            select(BetaSystemEvent)
            .where(
                BetaSystemEvent.organisation_id == organisation_id,
                BetaSystemEvent.id == owner_demo_base_manifest_id(organisation_id),
            )
            .with_for_update()
        )
        if tenant is None and owner_manifest is not None:
            if owner_manifest.actor_user_id is None:
                raise ValueError("Owner demo reset refused an unauthorised owner reservation.")
            membership = await session.get(
                OrganisationMembership,
                (organisation_id, owner_manifest.actor_user_id),
            )
            if membership is None or membership.status != "active" or membership.role != "admin":
                raise ValueError("Owner demo reset refused an unauthorised owner reservation.")
            tenant = TenantContext(
                organisation_id=organisation_id,
                user_id=owner_manifest.actor_user_id,
                role=cast(Role, membership.role),
            )
        owner_extensions: dict[str, UUID] | None = None
        if owner_manifest is not None and owner_manifest.event_type == OWNER_BASE_RESERVATION_EVENT:
            if tenant is None:
                raise ValueError("Owner demo reset refused an unauthorised owner reservation.")
            methodology_reservation = _validated_base_reservation(
                owner_manifest,
                organisation_id,
                tenant.user_id,
            )
            methodology_setting = await session.scalar(
                select(OrganisationMethodologySetting)
                .where(OrganisationMethodologySetting.organisation_id == organisation_id)
                .with_for_update()
            )
            _assert_recoverable_methodology_reservation_state(
                methodology_setting,
                tenant.user_id,
                methodology_reservation,
            )
        elif owner_manifest is not None and tenant is not None:
            owner_base = _validated_base_manifest(
                owner_manifest,
                organisation_id,
                tenant.user_id,
            )
            methodology_setting = await session.scalar(
                select(OrganisationMethodologySetting)
                .where(OrganisationMethodologySetting.organisation_id == organisation_id)
                .with_for_update()
            )
            _validate_completed_owner_methodology_setting(
                methodology_setting,
                tenant.user_id,
                owner_base,
                allow_missing=True,
            )
            owner_extensions = _validated_extension_manifest(
                owner_manifest,
                organisation_id,
                tenant.user_id,
            )
        company = await session.get(Company, company_id)
        reserved_rows = (
            await session.scalar(
                select(SellingProfileRevision.id).where(
                    SellingProfileRevision.organisation_id == organisation_id,
                    SellingProfileRevision.idempotency_key == OWNER_PROFILE_KEY,
                )
            ),
            await session.scalar(
                select(CreateValueModel.id).where(
                    CreateValueModel.organisation_id == organisation_id,
                    CreateValueModel.idempotency_key == OWNER_VALUE_MODEL_KEY,
                )
            ),
            await session.scalar(
                select(CreateBusinessCase.id).where(
                    CreateBusinessCase.organisation_id == organisation_id,
                    CreateBusinessCase.idempotency_key == OWNER_BUSINESS_CASE_KEY,
                )
            ),
            await session.scalar(
                select(CreatePresentation.id).where(
                    CreatePresentation.organisation_id == organisation_id,
                    CreatePresentation.idempotency_key == OWNER_PRESENTATION_KEY,
                )
            ),
            await session.scalar(
                select(CompanyKnowledgeSource.id).where(
                    CompanyKnowledgeSource.organisation_id == organisation_id,
                    CompanyKnowledgeSource.idempotency_key.in_((OWNER_KNOWLEDGE_WEBSITE_KEY, OWNER_KNOWLEDGE_TEXT_KEY)),
                )
            ),
        )
        reserved_rows_exist = any(row is not None for row in reserved_rows)
        deterministic_template = await session.scalar(
            select(CreateTemplate).where(
                CreateTemplate.organisation_id == organisation_id,
                CreateTemplate.id == owner_demo_template_id(organisation_id),
            )
        )
        owner_deal_room = await session.scalar(
            select(DealRoom.id).where(
                DealRoom.organisation_id == organisation_id,
                DealRoom.opportunity_id == opportunity_id,
            )
        )
        owner_handover = await session.scalar(
            select(ClosedWonHandover.id).where(
                ClosedWonHandover.organisation_id == organisation_id,
                ClosedWonHandover.opportunity_id == opportunity_id,
            )
        )
        owner_task = await session.scalar(
            select(Task)
            .where(
                Task.organisation_id == organisation_id,
                Task.id == owner_demo_task_id(organisation_id),
            )
            .with_for_update()
        )
        if company is None and (
            reserved_rows_exist
            or deterministic_template is not None
            or owner_deal_room is not None
            or owner_handover is not None
            or owner_task is not None
        ):
            raise ValueError(
                "Owner demo reset refused an identity/content collision without the exact synthetic Account owner marker."
            )

        if company is not None and (tenant is None or company.owner_user_id != tenant.user_id):
            raise ValueError("Owner demo reset refused: the exact synthetic Account owner could not be verified.")
        owner_revision: SellingProfileRevision | None = None
        owner_knowledge_sources: list[CompanyKnowledgeSource] = []
        value_model: CreateValueModel | None = None
        value_model_version: CreateValueModelVersion | None = None
        business_case: CreateBusinessCase | None = None
        business_case_version: CreateBusinessCaseVersion | None = None
        presentation: CreatePresentation | None = None
        validated_deal_room: tuple[DealRoom, Literal["create_only", "updated_draft", "published"]] | None = None
        validated_handover: ClosedWonHandover | None = None
        if tenant is not None:
            owner_revision = await _validated_owner_profile(session, tenant)
            owner_knowledge_sources = list(
                (
                    await session.scalars(
                        select(CompanyKnowledgeSource).where(
                            CompanyKnowledgeSource.organisation_id == organisation_id,
                            CompanyKnowledgeSource.idempotency_key.in_(
                                (OWNER_KNOWLEDGE_WEBSITE_KEY, OWNER_KNOWLEDGE_TEXT_KEY)
                            ),
                        )
                    )
                ).all()
            )
            if owner_knowledge_sources:
                by_key = {source.idempotency_key: source for source in owner_knowledge_sources}
                website_source = by_key.get(OWNER_KNOWLEDGE_WEBSITE_KEY)
                text_source = by_key.get(OWNER_KNOWLEDGE_TEXT_KEY)
                if (
                    len(owner_knowledge_sources) != 2
                    or website_source is None
                    or text_source is None
                    or website_source.source_kind != "website"
                    or website_source.label != "[DEMO] Oryntela website address"
                    or website_source.website_url != OWNER_KNOWLEDGE_URL
                    or website_source.source_text is not None
                    or website_source.content_fingerprint != hashlib.sha256(OWNER_KNOWLEDGE_URL.encode()).hexdigest()
                    or website_source.processing_status not in OWNER_KNOWLEDGE_PROCESSING_STATES
                    or website_source.created_by_user_id != tenant.user_id
                    or text_source.source_kind != "text_file"
                    or text_source.label != "[DEMO] Approved conversation product brief.txt"
                    or text_source.website_url is not None
                    or text_source.source_text != OWNER_KNOWLEDGE_TEXT
                    or text_source.content_fingerprint != hashlib.sha256(OWNER_KNOWLEDGE_TEXT.encode()).hexdigest()
                    or text_source.processing_status not in OWNER_KNOWLEDGE_PROCESSING_STATES
                    or text_source.created_by_user_id != tenant.user_id
                    or (
                        owner_extensions is not None
                        and (
                            owner_extensions["knowledge_website_source_id"] != website_source.id
                            or owner_extensions["knowledge_text_source_id"] != text_source.id
                        )
                    )
                ):
                    raise _owner_collision("Company Knowledge source")
            elif owner_extensions is not None:
                raise _owner_collision("Company Knowledge source")
            business_service = BusinessCaseService(session, tenant, active_settings)
            validated_model = await _validated_owner_value_model(session, tenant, business_service)
            if validated_model is not None:
                value_model, value_model_version = validated_model
                validated_case = await _validated_owner_business_case(
                    session,
                    tenant,
                    business_service,
                    company_id,
                    opportunity_id,
                    value_model,
                    value_model_version,
                )
                if validated_case is not None:
                    business_case, business_case_version = validated_case
            elif await session.scalar(
                select(CreateBusinessCase.id).where(
                    CreateBusinessCase.organisation_id == organisation_id,
                    CreateBusinessCase.idempotency_key == OWNER_BUSINESS_CASE_KEY,
                )
            ):
                raise _owner_collision("Business Case")

            demo_parent_links = [
                CreateBusinessCase.account_id == company_id,
                CreateBusinessCase.opportunity_id == opportunity_id,
            ]
            if value_model is not None:
                demo_parent_links.append(CreateBusinessCase.model_id == value_model.id)
            dependent_business_cases = select(CreateBusinessCase.id).where(
                CreateBusinessCase.organisation_id == organisation_id,
                or_(*demo_parent_links),
            )
            if business_case is not None:
                dependent_business_cases = dependent_business_cases.where(CreateBusinessCase.id != business_case.id)
            if await session.scalar(dependent_business_cases) is not None:
                raise ValueError(
                    "Owner demo reset refused because another Business Case references a "
                    "manifest-owned demo parent. Remove or relink that local record before resetting."
                )

            if business_case is not None and business_case_version is not None:
                presentation = await _validated_owner_presentation(
                    session,
                    tenant,
                    _create_service(session, tenant, active_settings),
                    company_id,
                    opportunity_id,
                    contact_id,
                    business_case,
                    business_case_version,
                    expected_template_id=None,
                    expected_template_version_id=None,
                    storage=storage,
                )
            elif await session.scalar(
                select(CreatePresentation.id).where(
                    CreatePresentation.organisation_id == organisation_id,
                    CreatePresentation.idempotency_key == OWNER_PRESENTATION_KEY,
                )
            ):
                raise _owner_collision("Presentation")

            presentation_parent_links = [
                CreatePresentation.account_id == company_id,
                CreatePresentation.opportunity_id == opportunity_id,
            ]
            if business_case is not None:
                presentation_parent_links.append(CreatePresentation.business_case_id == business_case.id)
            dependent_presentations = select(CreatePresentation.id).where(
                CreatePresentation.organisation_id == organisation_id,
                or_(*presentation_parent_links),
            )
            if presentation is not None:
                dependent_presentations = dependent_presentations.where(CreatePresentation.id != presentation.id)
            if await session.scalar(dependent_presentations) is not None:
                raise ValueError(
                    "Owner demo reset refused because another Presentation references a "
                    "manifest-owned demo Account, Opportunity or Business Case. Remove or relink "
                    "that local record before resetting."
                )

            if presentation is not None and business_case_version is not None:
                presentation_version_id = await session.scalar(
                    select(CreatePresentationVersion.id).where(
                        CreatePresentationVersion.organisation_id == organisation_id,
                        CreatePresentationVersion.presentation_id == presentation.id,
                    )
                )
                if presentation_version_id is None:
                    raise _owner_collision("Presentation")
                expected_room_content = _owner_deal_room_content(
                    organisation_id,
                    opportunity_id,
                    contact_id,
                    business_case_version.id,
                    presentation_version_id,
                )
                validated_deal_room = await _validated_owner_deal_room(
                    session,
                    tenant,
                    active_settings,
                    opportunity_id,
                    expected_room_content,
                    allow_owner_lifecycle=(
                        owner_extensions is not None and owner_extensions["deal_room_id"] == owner_deal_room
                    ),
                )
                handover_result = await _validated_owner_handover(
                    session,
                    tenant,
                    HandoverService(session, tenant, active_settings),
                    opportunity_id,
                    allow_owner_lifecycle=(
                        owner_extensions is not None and owner_extensions["handover_id"] == owner_handover
                    ),
                )
                if handover_result is not None:
                    validated_handover = handover_result[0]
        if owner_deal_room is not None and validated_deal_room is None:
            raise _owner_collision("Deal Room")
        if owner_handover is not None and validated_handover is None:
            raise _owner_collision("Closed-Won Handover")
        if owner_task is not None:
            task_due_at = owner_task.due_at
            if task_due_at is not None and task_due_at.utcoffset() is None:
                task_due_at = task_due_at.replace(tzinfo=UTC)
            if (
                tenant is None
                or (owner_extensions is not None and owner_extensions["task_id"] != owner_task.id)
                or owner_task.company_id != company_id
                or owner_task.contact_id != contact_id
                or owner_task.opportunity_id != opportunity_id
                or owner_task.title != "[DEMO] Prepare the Northstar security workshop"
                or owner_task.description
                != "Synthetic owner-demo follow-up. Confirm the reviewed security summary and workshop agenda."
                or owner_task.status not in {"open", "in_progress", "completed", "cancelled"}
                or owner_task.priority != "high"
                or task_due_at != datetime(2026, 11, 4, 2, 0, tzinfo=UTC)
                or owner_task.assigned_user_id != tenant.user_id
                or owner_task.created_by_user_id != tenant.user_id
            ):
                raise _owner_collision("Task")

        presentations = [presentation] if presentation is not None else []
        presentation_ids = [item.id for item in presentations]
        candidate_template_ids = {owner_demo_template_id(organisation_id)}
        candidate_template_ids.update(item.template_id for item in presentations)
        templates = list(
            (
                await session.scalars(
                    select(CreateTemplate).where(
                        CreateTemplate.organisation_id == organisation_id,
                        CreateTemplate.id.in_(candidate_template_ids),
                    )
                )
            ).all()
        )
        template_ids: list[UUID] = []
        template_storage_keys: list[str] = []
        expected_template = _template_bytes()
        expected_checksum = hashlib.sha256(expected_template).hexdigest()
        for template in templates:
            versions = list(
                (
                    await session.scalars(
                        select(CreateTemplateVersion).where(
                            CreateTemplateVersion.organisation_id == organisation_id,
                            CreateTemplateVersion.template_id == template.id,
                        )
                    )
                ).all()
            )
            if (
                tenant is None
                or template.name != OWNER_TEMPLATE_NAME
                or template.state != "active"
                or template.created_by_user_id != tenant.user_id
            ):
                raise ValueError("Owner demo reset refused a template identity/content collision.")
            if template.id != owner_demo_template_id(organisation_id) and not versions:
                raise ValueError("Owner demo reset refused an unverified legacy template marker.")
            if not versions:
                raise ValueError("Owner demo reset refused an unverified template marker.")
            for version in versions:
                if (
                    version.display_filename != OWNER_TEMPLATE_FILE_NAME
                    or version.mime_type != PPTX_MIME_TYPE
                    or version.uploaded_by_user_id != tenant.user_id
                    or version.authority_attested_by_user_id != tenant.user_id
                ):
                    raise ValueError("Owner demo reset refused a template identity/content collision.")
                try:
                    stored = await storage.read(version.storage_key)
                    stored_checksum = hashlib.sha256(stored).hexdigest()
                    canonical_checksum = hashlib.sha256(_canonicalise_pptx_archive(stored)).hexdigest()
                except (OSError, VisualStorageError, zipfile.BadZipFile) as exc:
                    raise ValueError("Owner demo reset could not verify its captured local template.") from exc
                if (
                    len(stored) != version.byte_size
                    or stored_checksum != version.checksum_sha256
                    or canonical_checksum != expected_checksum
                ):
                    raise ValueError("Owner demo reset refused a template identity/content collision.")
            template_ids.append(template.id)
            template_storage_keys.extend(version.storage_key for version in versions)
        presentation_versions = list(
            (
                await session.scalars(
                    select(CreatePresentationVersion).where(
                        CreatePresentationVersion.organisation_id == organisation_id,
                        CreatePresentationVersion.presentation_id.in_(presentation_ids),
                    )
                )
            ).all()
        )
        presentation_storage_keys: list[str] = []
        for presentation_version in presentation_versions:
            if presentation_version.pptx_storage_key is None:
                continue
            try:
                stored = await storage.read(presentation_version.pptx_storage_key)
            except (OSError, VisualStorageError) as exc:
                raise ValueError("Owner demo reset could not verify its captured local presentation.") from exc
            if (
                presentation_version.byte_size is None
                or presentation_version.checksum_sha256 is None
                or len(stored) != presentation_version.byte_size
                or hashlib.sha256(stored).hexdigest() != presentation_version.checksum_sha256
            ):
                raise ValueError("Owner demo reset refused a presentation identity/content collision.")
            presentation_storage_keys.append(presentation_version.pptx_storage_key)
        # A same-named template without the deterministic identity or an exact
        # owner-presentation link is legitimate user data and remains untouched.
        owner_revision_ids = [owner_revision.id] if owner_revision is not None else []
        owner_profile_ids = [owner_revision.profile_id] if owner_revision is not None else []
        owner_knowledge_source_ids = [source.id for source in owner_knowledge_sources]
        captured_keys = tuple(sorted({*template_storage_keys, *presentation_storage_keys}))
        _validated_cleanup_storage_keys(organisation_id, list(captured_keys))
        cleanup_keys = tuple(sorted({*pending_cleanup_keys, *captured_keys}))
        await session.execute(
            delete(ClosedWonHandover).where(
                ClosedWonHandover.organisation_id == organisation_id,
                ClosedWonHandover.id == (validated_handover.id if validated_handover is not None else None),
            )
        )
        await session.execute(
            delete(DealRoom).where(
                DealRoom.organisation_id == organisation_id,
                DealRoom.id == (validated_deal_room[0].id if validated_deal_room is not None else None),
            )
        )
        await session.execute(
            delete(Task).where(
                Task.organisation_id == organisation_id,
                Task.id == (owner_task.id if owner_task is not None else None),
            )
        )
        await session.execute(
            delete(CreatePresentation).where(
                CreatePresentation.organisation_id == organisation_id,
                CreatePresentation.id.in_(presentation_ids),
            )
        )
        await session.execute(
            delete(CreateBusinessCase).where(
                CreateBusinessCase.organisation_id == organisation_id,
                CreateBusinessCase.id == (business_case.id if business_case is not None else None),
            )
        )
        await session.execute(
            delete(CreateValueModel).where(
                CreateValueModel.organisation_id == organisation_id,
                CreateValueModel.id == (value_model.id if value_model is not None else None),
            )
        )
        await session.execute(
            delete(CreateTemplate).where(
                CreateTemplate.organisation_id == organisation_id,
                CreateTemplate.id.in_(template_ids),
            )
        )
        await session.execute(
            delete(BetaSystemEvent).where(
                BetaSystemEvent.organisation_id == organisation_id,
                BetaSystemEvent.subject_id.in_((*owner_revision_ids, *owner_knowledge_source_ids)),
            )
        )
        await session.execute(
            delete(CompanyKnowledgeSource).where(
                CompanyKnowledgeSource.organisation_id == organisation_id,
                CompanyKnowledgeSource.id.in_(owner_knowledge_source_ids),
            )
        )
        await session.execute(
            delete(SellingProfileRevision).where(
                SellingProfileRevision.organisation_id == organisation_id,
                SellingProfileRevision.id.in_(owner_revision_ids),
            )
        )
        for profile_id in owner_profile_ids:
            remaining = await session.scalar(
                select(SellingProfileRevision.id).where(
                    SellingProfileRevision.organisation_id == organisation_id,
                    SellingProfileRevision.profile_id == profile_id,
                )
            )
            if remaining is None:
                await session.execute(
                    delete(SellingProfile).where(
                        SellingProfile.organisation_id == organisation_id,
                        SellingProfile.id == profile_id,
                    )
                )
        base_cleanup_keys: Sequence[str] = ()
        if database_finalize is not None:
            # The full demo reset supplies its base-row deletion here so every
            # database mutation shares this transaction. A restrictive local
            # dependent therefore rolls back the owner extensions and cleanup
            # marker before any captured owner-demo object is removed.
            base_cleanup_keys = await database_finalize(session)
            _validated_cleanup_storage_keys(organisation_id, list(base_cleanup_keys))
        cleanup_keys = tuple(sorted({*cleanup_keys, *base_cleanup_keys}))
        if cleanup_keys:
            metadata: dict[str, object] = {
                "schemaVersion": OWNER_RESET_CLEANUP_SCHEMA_VERSION,
                "storageKeys": list(cleanup_keys),
            }
            if cleanup_event is None:
                session.add(
                    BetaSystemEvent(
                        id=cleanup_event_id,
                        organisation_id=organisation_id,
                        actor_user_id=tenant.user_id if tenant is not None else None,
                        event_type=OWNER_RESET_CLEANUP_EVENT,
                        subject_id=cleanup_event_id,
                        metadata_json=metadata,
                    )
                )
            else:
                cleanup_event.metadata_json = metadata

    cleanup_errors: list[Exception] = []
    for storage_key in cleanup_keys:
        try:
            await storage.delete(storage_key)
        except (OSError, VisualStorageError) as exc:
            cleanup_errors.append(exc)
    if cleanup_errors:
        raise VisualStorageError(
            f"Owner demo database reset committed, but local object cleanup failed for {len(cleanup_errors)} object(s)."
        ) from cleanup_errors[0]
    if not cleanup_keys:
        return
    async with session_factory() as session, session.begin():
        await set_tenant_database_context(session, organisation_id)
        if session.get_bind().dialect.name == "postgresql":
            await session.execute(text("SELECT set_config('app.beta_maintenance', 'approved', true)"))
        cleanup_event = await session.scalar(
            select(BetaSystemEvent)
            .where(
                BetaSystemEvent.organisation_id == organisation_id,
                BetaSystemEvent.id == cleanup_event_id,
            )
            .with_for_update()
        )
        if cleanup_event is None:
            raise VisualStorageError("Owner demo local object cleanup marker disappeared before acknowledgement.")
        current_keys = _cleanup_keys_from_event(cleanup_event, organisation_id)
        remaining_keys = tuple(key for key in current_keys if key not in cleanup_keys)
        if remaining_keys:
            cleanup_event.metadata_json = {
                "schemaVersion": OWNER_RESET_CLEANUP_SCHEMA_VERSION,
                "storageKeys": list(remaining_keys),
            }
        else:
            await session.delete(cleanup_event)
